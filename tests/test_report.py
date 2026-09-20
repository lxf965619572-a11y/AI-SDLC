"""软件测评报告装配（report_exporter）的测试：结论 / 证据清单 / 问题单 / 偏差单 / 导出。

第一性原理：测评报告是交付件的「结论出口」，它自己不下判断，只把已落库的工具
结论折叠成人能审、能归档的文档。所以这里钉的是三条不能错的契约：
  1. 结论保守——静态与执行任一判据缺失或不通过，都不得给出「通过」；
  2. 证据可回溯——证据清单里的 sha256 必须与同步输入逐字节对得上，
     这是「报告所述代码 == 实际执行代码」的唯一凭据；
  3. 问题归零留痕——闭环过程中失败过的轮次要登记为问题报告单，
     修好了标「已闭环」，没修好标「转人工」，绝不因为最终绿了就抹掉历史。

report_service 依赖数据库，本文件只测 exporters.report_exporter 的纯装配逻辑；
另含一条「真实 executor 绿run → 报告」的集成测试，防止两层产物形状漂移。

零第三方依赖（openpyxl / python-docx 已在 requirements 内）：
    .venv\\Scripts\\python.exe tests\\test_report.py
"""
import hashlib
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from exporters import report_exporter as REX
from tests.fixtures import comm_sample as SAMPLE


def _sha(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


# ---------- 文档阶段 meta（追溯摘要用，形状与 test_trace_verify 一致）----------
REQ = {"functional_requirements": [
    {"id": "FR-001", "desc": "成帧并登记重传窗口", "priority": "P0",
     "derived_from": ["OBJ-001"]},
    {"id": "FR-002", "desc": "CRC-16 校验", "priority": "P0",
     "derived_from": ["RULE-001"]},
]}
HLD = {"modules": ["comm"], "derived_from": {"comm": ["FR-001", "FR-002"]}}
LLD = {"functions": [{"name": "comm_pack"}, {"name": "comm_crc16"}],
       "derived_from": {"comm_pack": ["FR-001"], "comm_crc16": ["FR-002"]}}
TC = {"testcases": [{"id": "TC-001", "fr_ids": ["FR-002"]},
                    {"id": "TC-003", "fr_ids": ["FR-001"]}]}
CODE = {"files": ["src/comm.c", "include/comm.h"],
        "derived_from": {"comm_pack": ["FR-001"], "comm_crc16": ["FR-002"]}}
TEST_IMPL = {"files": ["tests/test_comm.c"],
             "cases": [{"id": "TC-001", "fn": "comm_crc16", "fr_ids": ["FR-002"]},
                       {"id": "TC-003", "fn": "comm_pack", "fr_ids": ["FR-001"]}]}

STATIC_OK = {"ok": True, "required": 0, "advisory": 0, "total": 7, "by_rule": {},
             "violations": [],
             "functions": [{"name": "comm_pack", "file": "src/comm.c"},
                           {"name": "comm_crc16", "file": "src/comm.c"}],
             "thresholds": {"complexity_max": 10}}


def _static_fail(waivable=True):
    return {"ok": False, "required": 1, "advisory": 0, "total": 7,
            "by_rule": {"WB-C-005": 1},
            "violations": [{"rule_id": "WB-C-005", "title": "圈复杂度上限",
                            "severity": "required", "file": "src/comm.c", "line": 42,
                            "func": "comm_pack", "message": "圈复杂度 12 > 10",
                            "waivable": waivable, "standard_ref": "GJB 8114"}],
            "functions": [{"name": "comm_pack", "file": "src/comm.c"}],
            "thresholds": {"complexity_max": 10}}


# ---------- 执行 meta：证据清单的 sha256 用样例内容真算，便于独立校验 ----------
MANIFEST = {
    "src/comm.c": {"sha256": _sha(SAMPLE.COMM_C),
                   "bytes": len(SAMPLE.COMM_C.encode("utf-8"))},
    "include/comm.h": {"sha256": _sha(SAMPLE.COMM_H),
                       "bytes": len(SAMPLE.COMM_H.encode("utf-8"))},
    "tests/test_comm.c": {"sha256": _sha(SAMPLE.TEST_COMM_C),
                          "bytes": len(SAMPLE.TEST_COMM_C.encode("utf-8"))},
}


def _exec_meta(**over):
    base = {
        "ok": True, "skipped": False, "reason": "",
        "verdict": {"build": "ok", "tests": "ok", "coverage": "ok"},
        "env": {"uname": "Linux 4.15.0-156-generic x86_64",
                "gcc": "gcc (Ubuntu 7.5.0) 7.5.0", "gcov": "gcov (Ubuntu 7.5.0) 7.5.0",
                "make": "GNU Make 4.1", "cores": "4"},
        "transport": "ssh", "workdir": "/home/lixf/wb_verify/p1/v1t1",
        "duration_s": 3.2,
        "commands": [{"cmd": "sh build.sh", "cwd": "/home/lixf/wb_verify/p1/v1t1",
                      "exit_code": 0}],
        "build": {"ok": True, "warning_count": 0, "errors": []},
        "tests": {"expected": ["TC-001", "TC-003"], "total": 2, "passed": 2,
                  "failed": 0, "missing": 0, "all_pass": True, "consistent": True,
                  "crashed": False, "run_exit": 0,
                  "results": [{"id": "TC-001", "status": "pass", "status_text": "通过",
                               "detail": ""},
                              {"id": "TC-003", "status": "pass", "status_text": "通过",
                               "detail": ""}]},
        "coverage": {
            "totals": {"line_pct": 100.0, "lines_hit": 82, "lines_total": 82,
                       "branch_pct": 100.0, "branch_taken": 26, "branch_total": 26},
            "functions": [
                {"kind": "function", "name": "comm_pack", "file": "src/comm.c",
                 "lines_pct": 100.0, "branch_effective": 100.0, "branch_total": 6,
                 "never_executed": False},
                {"kind": "function", "name": "comm_crc16", "file": "src/comm.c",
                 "lines_pct": 100.0, "branch_effective": 100.0, "branch_total": 4,
                 "never_executed": False}],
            "files": [], "below_branch": [], "below_line": [], "untested_functions": [],
            "thresholds": {"branch_min": 80.0, "line_min": 0.0}, "ok": True},
        "manifest": MANIFEST,
        "evidence": "data/verify/p1/v1t1/exec.log",
        "thresholds": {"branch_min": 80.0, "line_min": 0.0},
    }
    base.update(over)
    return base


EXEC_GREEN = _exec_meta()


def make_arts(static=STATIC_OK, exec_meta=EXEC_GREEN, static_status="approved",
              with_exec=True):
    arts = {
        "requirement": {"markdown": "# SRS", "meta": REQ, "version": 1, "status": "approved"},
        "hld": {"markdown": "# HLD", "meta": HLD, "version": 1, "status": "approved"},
        "lld": {"markdown": "# LLD", "meta": LLD, "version": 1, "status": "approved"},
        "testcase": {"markdown": "# TC", "meta": TC, "version": 1, "status": "approved"},
        "code": {"markdown": "# CODE", "meta": CODE, "version": 1, "status": "approved"},
        "static": {"markdown": "# STATIC", "meta": static, "version": 1,
                   "status": static_status},
        "test_impl": {"markdown": "# TI", "meta": TEST_IMPL, "version": 1,
                      "status": "approved"},
    }
    if with_exec:
        arts["exec"] = {"markdown": "# EXEC", "meta": exec_meta, "version": 1,
                        "status": "approved"}
    return arts


def assemble(arts, **kw):
    kw.setdefault("coverage_min", 80.0)
    return REX.assemble_report(arts, "通信协议样例模块", **kw)


# ================= 测评结论 =================
def test_conclusion_pass_when_static_and_exec_ok():
    data = assemble(make_arts())
    assert data["conclusion"]["pass"] is True
    assert "全部通过" in data["conclusion"]["text"]
    assert data["conclusion"]["reasons"] == []


def test_conclusion_fails_when_exec_missing():
    data = assemble(make_arts(with_exec=False))
    assert data["conclusion"]["pass"] is False
    assert any("未产出验证执行结论" in r for r in data["conclusion"]["reasons"])


def test_conclusion_fails_when_exec_skipped():
    em = _exec_meta(ok=False, skipped=True, reason="未配置验证机",
                    verdict={"build": "skipped", "tests": "skipped", "coverage": "skipped"})
    data = assemble(make_arts(exec_meta=em))
    assert data["conclusion"]["pass"] is False
    assert any("验证未执行" in r for r in data["conclusion"]["reasons"])


def test_conclusion_fails_when_exec_not_ok():
    em = _exec_meta(ok=False, reason="1 条用例失败",
                    verdict={"build": "ok", "tests": "fail", "coverage": "ok"})
    data = assemble(make_arts(exec_meta=em))
    assert data["conclusion"]["pass"] is False
    assert any("验证执行未通过" in r for r in data["conclusion"]["reasons"])


def test_conclusion_fails_when_static_required_violation():
    data = assemble(make_arts(static=_static_fail()))
    assert data["conclusion"]["pass"] is False
    assert any("必查项违规 1 条" in r for r in data["conclusion"]["reasons"])


def test_conclusion_fails_when_harness_inconsistent():
    """全绿但测试桩自报数对不上：结论不可信，必须拦下。"""
    tests = dict(EXEC_GREEN["tests"], consistent=False, reported_total=99)
    em = _exec_meta(tests=tests)
    data = assemble(make_arts(exec_meta=em))
    assert data["conclusion"]["pass"] is False
    assert any("不一致" in r for r in data["conclusion"]["reasons"])


# ================= 证据清单（manifest 哈希）=================
def test_evidence_lists_synced_inputs_with_real_hashes():
    data = assemble(make_arts())
    synced = [e for e in data["evidence"] if e["kind"] == "同步输入"]
    by_path = {e["path"]: e for e in synced}
    assert set(by_path) == set(MANIFEST)
    # sha256 与独立计算逐字节一致——这是「报告所述代码==实际执行代码」的凭据
    assert by_path["src/comm.c"]["sha256"] == _sha(SAMPLE.COMM_C)
    assert by_path["src/comm.c"]["bytes"] == len(SAMPLE.COMM_C.encode("utf-8"))


def test_evidence_includes_raw_output_archive():
    data = assemble(make_arts())
    arch = [e for e in data["evidence"] if e["kind"] == "原始输出归档"]
    assert len(arch) == 1
    assert arch[0]["path"] == "data/verify/p1/v1t1/exec.log"


def test_evidence_no_archive_when_exec_has_none():
    em = _exec_meta(evidence="")
    data = assemble(make_arts(exec_meta=em))
    assert all(e["kind"] != "原始输出归档" for e in data["evidence"])


def test_evidence_legacy_string_manifest_compat():
    """早期产物 manifest 只存哈希字符串：报告要兼容，bytes 缺省为 None。"""
    em = _exec_meta(manifest={"src/comm.c": _sha(SAMPLE.COMM_C)})
    data = assemble(make_arts(exec_meta=em))
    synced = [e for e in data["evidence"] if e["kind"] == "同步输入"]
    assert len(synced) == 1
    assert synced[0]["sha256"] == _sha(SAMPLE.COMM_C)
    assert synced[0]["bytes"] is None


# ================= 问题报告单（历史留痕）=================
def test_problem_registered_and_resolved_when_later_round_ok():
    history = {"exec": [
        {"version": 1, "status": "rejected",
         "meta": {"ok": False, "reason": "1 条用例失败",
                  "attribution": {"decision": "fix_code", "reason": "CRC 初值错",
                                  "source": "agent"}}},
        {"version": 2, "status": "approved", "meta": {"ok": True}},
    ]}
    data = assemble(make_arts(), history=history)
    probs = data["problems"]
    assert len(probs) == 1
    p = probs[0]
    assert p["id"] == "PR-001" and p["stage"] == "exec"
    assert p["resolved"] is True
    assert "已闭环" in p["disposition"]
    assert p["decision_label"] == "被测代码缺陷"
    assert p["attribution"] == "CRC 初值错"


def test_problem_unresolved_when_no_later_ok():
    history = {"exec": [
        {"version": 1, "status": "rejected", "meta": {"ok": False, "reason": "覆盖率不足"}}]}
    data = assemble(make_arts(), history=history)
    p = data["problems"][0]
    assert p["resolved"] is False
    assert "转人工" in p["disposition"]


def test_problem_from_static_round():
    history = {"static": [
        {"version": 1, "status": "rejected",
         "meta": {"ok": False, "required": 1,
                  "violations": [{"rule_id": "WB-C-005", "severity": "required"}]}}]}
    data = assemble(make_arts(static=_static_fail()), history=history)
    probs = [p for p in data["problems"] if p["stage"] == "static"]
    assert len(probs) == 1
    assert "WB-C-005" in probs[0]["detail"]


def test_no_problem_when_all_rounds_ok():
    data = assemble(make_arts())
    assert data["problems"] == []


# ================= 偏差单 =================
def test_deviation_waivable_approved():
    data = assemble(make_arts(static=_static_fail(waivable=True),
                              static_status="approved"))
    dev = data["deviations"][0]
    assert dev["id"] == "DEV-001" and dev["rule_id"] == "WB-C-005"
    assert dev["waivable"] is True
    assert dev["disposition"] == "人工批准放行"


def test_deviation_waivable_not_approved():
    data = assemble(make_arts(static=_static_fail(waivable=True),
                              static_status="pending_review"))
    assert data["deviations"][0]["disposition"] == "待人工裁决"


def test_deviation_non_waivable_must_fix():
    data = assemble(make_arts(static=_static_fail(waivable=False),
                              static_status="approved"))
    assert data["deviations"][0]["disposition"] == "不可偏差，须整改"


def test_no_deviation_when_static_ok():
    data = assemble(make_arts())
    assert data["deviations"] == []


# ================= 追溯摘要 =================
def test_trace_summary_reuses_core_trace():
    data = assemble(make_arts())
    tr = data["trace"]
    assert tr["fr_total"] == 2 and tr["p0_total"] == 2
    assert tr["status_counts"]            # 非空
    assert isinstance(tr["rows"], list) and tr["rows"]


# ================= Markdown 渲染 =================
def test_markdown_renders_all_sections():
    md = REX.report_markdown(assemble(make_arts()))
    for head in ("# 软件测评报告", "## 3 测评综述", "## 4 测评结果",
                 "## 5 需求追溯摘要", "## 8 测评结论", "## 9 证据清单"):
        assert head in md, head
    assert "src/comm.c" in md            # 证据清单里有同步输入
    assert "通过" in md


def test_markdown_conclusion_not_pass_when_exec_skipped():
    em = _exec_meta(ok=False, skipped=True, reason="未配置验证机",
                    verdict={"build": "skipped", "tests": "skipped", "coverage": "skipped"})
    md = REX.report_markdown(assemble(make_arts(exec_meta=em)))
    assert "## 8 测评结论" in md
    assert "未通过" in md


# ================= 文件导出 =================
def test_export_excel_has_sheets_and_evidence_hash():
    from openpyxl import load_workbook
    data = assemble(make_arts())
    with tempfile.TemporaryDirectory() as tmp:
        path = str(Path(tmp) / "report.xlsx")
        REX.export_report_excel(data, path)
        assert os.path.getsize(path) > 0
        wb = load_workbook(path)
        for name in ("测评结论", "用例结果", "覆盖率", "静态检查", "问题与偏差", "证据清单"):
            assert name in wb.sheetnames, name
        ws = wb["证据清单"]
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
        hit = [r for r in rows if r[1] == "src/comm.c"]
        assert hit and hit[0][3] == _sha(SAMPLE.COMM_C)


def test_export_docx_creates_file():
    data = assemble(make_arts())
    with tempfile.TemporaryDirectory() as tmp:
        path = str(Path(tmp) / "report.docx")
        REX.export_report_docx(data, path)
        assert os.path.getsize(path) > 0


# ================= 集成：真实 executor 绿run → 报告 =================
def test_report_consumes_real_executor_outcome():
    """防形状漂移：报告必须能直接吃下 executor.run_verification 的真实产物。"""
    from tests import test_executor as TEX
    outcome = TEX._green_outcome()
    assert outcome["ok"] is True
    # 按 pipeline.nodes._slim_exec 的口径精简入库
    drop = ("raw_log", "run_section", "coverage_section", "synced_files")
    slim = {k: v for k, v in outcome.items() if k not in drop}
    b = outcome["build"]
    slim["build"] = {"ok": b["ok"], "warning_count": len(b.get("warnings") or []),
                     "errors": b.get("errors") or [], "log_tail": ""}
    slim["evidence"] = "data/verify/p1/v1t1/exec.log"
    arts = make_arts(exec_meta=slim)
    data = assemble(arts)
    assert data["conclusion"]["pass"] is True
    synced = {e["path"]: e for e in data["evidence"] if e["kind"] == "同步输入"}
    assert "src/comm.c" in synced
    # 真实 manifest 的哈希同样要能对上样例内容
    assert synced["src/comm.c"]["sha256"] == _sha(SAMPLE.COMM_C)


def _main():
    tests = [(n, f) for n, f in sorted(globals().items())
             if n.startswith("test_") and callable(f)]
    failed = []
    for name, fn in tests:
        try:
            fn()
            print("PASS  " + name)
        except Exception as e:
            failed.append(name)
            print("FAIL  " + name + ": " + type(e).__name__ + ": " + str(e))
    print("\n%d/%d passed" % (len(tests) - len(failed), len(tests)))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(_main())
