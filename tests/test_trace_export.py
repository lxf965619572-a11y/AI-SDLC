"""需求追溯矩阵 Excel 导出测试：写盘 → openpyxl 重新打开 → 校验内容。

导出是交付件，光看函数返回值不够，必须验证落到磁盘上的表真的能读回来、
表头/状态/覆盖数字都对得上。依赖 openpyxl（requirements 已含）。

用法：.venv\\Scripts\\python.exe tests\\test_trace_export.py
"""
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

from core import trace
from exporters.trace_exporter import HEADERS, export_trace_excel

SRS = {"functional_requirements": [
    {"id": "FR-001", "desc": "初始化通信", "priority": "P0", "derived_from": ["OBJ-001"]},
    {"id": "FR-002", "desc": "数据重传", "priority": "P0", "derived_from": ["RULE-002"]},
]}
HLD = {"modules": ["comm", "retry"],
       "derived_from": {"comm": ["FR-001"], "retry": ["FR-002"]}}
LLD = {"derived_from": {"comm_init": ["comm"], "retry_run": ["retry"]}}
TC = {"testcases": [{"id": "TC-1", "fr_ids": ["FR-001"]}]}

SOURCE_NAMES = {"OBJ-001": "通信对象", "RULE-002": "重传规则"}


def _matrix(with_sources=True):
    """模拟 services.trace_service.build_project_matrix 的返回结构。"""
    data = trace.build_matrix(SRS, HLD, LLD, TC)
    data["versions"] = {"requirement": 1, "hld": 1, "lld": 1, "testcase": 1}
    data["sources"] = {k: {"kind": k[:3], "name": v}
                       for k, v in SOURCE_NAMES.items()}
    for row in data["rows"]:
        if with_sources:
            row["source_names"] = [SOURCE_NAMES.get(s) or s for s in row["sources"]]
    return data


def _export(data, name="trace.xlsx"):
    tmp = Path(tempfile.mkdtemp(prefix="wb_trace_"))
    try:
        path = export_trace_excel(data, str(tmp / name), project_name="演示项目")
        return load_workbook(path), tmp
    except Exception:
        shutil.rmtree(tmp, ignore_errors=True)
        raise


def _overview(wb):
    ws = wb["覆盖概览"]
    return {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
            for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value}


def _by_header(wb, sheet="追溯矩阵"):
    """按表头名取每行的值：矩阵是要加列的，写死列号会让测试变成维护负担。"""
    ws = wb[sheet]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    out = {}
    for r in range(2, ws.max_row + 1):
        key = ws.cell(row=r, column=1).value
        out[key] = {h: ws.cell(row=r, column=i + 1).value
                    for i, h in enumerate(headers)}
    return out


def test_export_writes_two_sheets_with_headers():
    wb, tmp = _export(_matrix())
    try:
        assert wb.sheetnames == ["追溯矩阵", "覆盖概览"]
        ws = wb["追溯矩阵"]
        assert [ws.cell(row=1, column=c).value for c in range(1, len(HEADERS) + 1)] == HEADERS
        assert ws.freeze_panes == "A2"
        assert ws.max_row == 3                      # 表头 + 2 条需求
    finally:
        wb.close()
        shutil.rmtree(tmp, ignore_errors=True)


def test_export_rows_carry_chain_and_status():
    wb, tmp = _export(_matrix())
    try:
        rows = _by_header(wb)
        assert rows["FR-001"]["素材来源"] == "OBJ-001 通信对象"   # 编号带上了名称
        assert rows["FR-001"]["概要设计"] == "comm"
        assert rows["FR-001"]["详细设计"] == "comm_init"
        assert rows["FR-001"]["测试用例"] == "TC-1"
        assert rows["FR-001"]["链路状态"] == "贯通"
        # 没跑代码验证闭环：新增四列一律「未产出」，不能凭空编数字
        for col in ("代码单元", "静态检查", "执行结果", "分支覆盖"):
            assert rows["FR-001"][col] == "未产出", col
        # FR-002 没被任何用例覆盖：状态必须点名缺哪一环
        assert not rows["FR-002"]["测试用例"]
        assert rows["FR-002"]["链路状态"] == "待补全：测试用例"
    finally:
        wb.close()
        shutil.rmtree(tmp, ignore_errors=True)


def test_export_overview_reports_p0_coverage():
    wb, tmp = _export(_matrix())
    try:
        ov = _overview(wb)
        assert ov["项目"] == "演示项目"
        assert ov["需求总数"] == 2 and ov["其中 P0"] == 2
        assert ov["P0 覆盖 · 概要设计"] == "2/2"
        assert ov["P0 覆盖 · 测试用例"] == "1/2（未覆盖：FR-002）"
        assert ov["未关联需求的用例"] == "无"
        assert ov["链路状态分布"] == "待补全 1 · 贯通 1"
        assert "需求规格 v1" in ov["产物版本"]
    finally:
        wb.close()
        shutil.rmtree(tmp, ignore_errors=True)


def test_export_without_source_names_falls_back_to_ids():
    # 直接用 build_matrix 的结果（没有 source_names 映射）也不能崩，退回纯编号
    data = trace.build_matrix(SRS, HLD, LLD, TC)
    wb, tmp = _export(data)
    try:
        rows = _by_header(wb)
        assert rows["FR-001"]["素材来源"] == "OBJ-001"
    finally:
        wb.close()
        shutil.rmtree(tmp, ignore_errors=True)


def test_export_empty_matrix_still_writes_valid_file():
    wb, tmp = _export(trace.build_matrix(None, None, None, None))
    try:
        ws = wb["追溯矩阵"]
        assert [ws.cell(row=1, column=c).value for c in range(1, len(HEADERS) + 1)] == HEADERS
        assert ws.max_row == 1
        ov = _overview(wb)
        assert ov["需求总数"] == 0 and ov["链路状态分布"] == "-"
    finally:
        wb.close()
        shutil.rmtree(tmp, ignore_errors=True)


def _main():
    tests = [(n, f) for n, f in sorted(globals().items())
             if n.startswith("test_") and callable(f)]
    failed = []
    for name, fn in tests:
        try:
            fn()
            print("PASS  " + name)
        except Exception as e:
            failed.append(name)
            print("FAIL  " + name + ": " + type(e).__name__ + ": " + str(e))
    print("\n%d/%d passed" % (len(tests) - len(failed), len(tests)))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(_main())
