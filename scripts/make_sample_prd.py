"""生成一份示例 PRD（docx），用于端到端测试流水线。"""
from docx import Document
from docx.shared import Pt

doc = Document()
doc.add_heading("电商交易系统 PRD", level=1)

doc.add_heading("1. 项目背景", level=2)
doc.add_paragraph(
    "公司计划建设自营电商平台，支持用户在线浏览商品、下单购买、支付与退款。"
    "一期聚焦核心交易链路，暂不涉及营销与物流模块。")

doc.add_heading("2. 用户角色", level=2)
doc.add_paragraph("普通用户：可注册、登录、浏览、下单、支付、申请退款。")
doc.add_paragraph("客服人员：可审核退款申请。")

doc.add_heading("3. 功能需求", level=2)
doc.add_heading("3.1 注册与登录", level=3)
doc.add_paragraph("用户使用手机号注册，需短信验证码校验。密码长度 8-20 位。")
doc.add_heading("3.2 商品浏览", level=3)
doc.add_paragraph("支持按分类浏览与关键词搜索，商品详情页展示价格、库存与规格。")
doc.add_heading("3.3 下单与支付", level=3)
doc.add_paragraph(
    "用户将商品加入购物车后提交订单。订单金额必须大于 0 且不超过 50000 元。"
    "下单成功后立即扣减库存。支付支持微信与支付宝。支付超时 30 分钟订单自动取消并释放库存。")
doc.add_heading("3.4 退款", level=3)
doc.add_paragraph(
    "用户可对已支付订单发起退款申请，客服审核通过后原路退回，并通知用户。")

doc.add_heading("4. 会员规则", level=2)
doc.add_paragraph("金卡会员享受 95 折，银卡会员享受 98 折，折扣在提交订单时计算。")

doc.add_heading("5. 非功能需求", level=2)
doc.add_paragraph("下单接口响应时间 P99 < 500ms；系统可用性 99.9%。")

doc.add_paragraph("")
tbl = doc.add_table(rows=3, cols=3)
tbl.style = "Table Grid"
for i, row in enumerate([
    ["状态", "说明", "可流转到"],
    ["CREATED", "订单已创建待支付", "PAID / CANCELLED"],
    ["PAID", "已支付", "REFUNDING"],
]):
    for j, v in enumerate(row):
        tbl.rows[i].cells[j].text = v

out = "sample_prd.docx"
doc.save(out)
print("saved", out)
