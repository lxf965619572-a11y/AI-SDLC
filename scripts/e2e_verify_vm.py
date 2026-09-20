"""端到端验收：在真实验证机（Ubuntu VM）上跑通「生成 + 验证闭环」的判据层。

为什么单独一个脚本：流水线全链路（requirement→…→report）要过 6 道人工评审门、
要烧真实模型 token，属于「手动验收」。但整套系统真正下结论、真正当证据用的，是
code→static→test_impl→exec→report 这条判据链——它不依赖模型，只依赖工具链事实。
本脚本把这条链在真实 VM 上跑一遍，一条命令即可复算，无需 Web UI、无需 LLM：

  1. 探测验证机工具链（gcc/gcov/make/tar），不可达即快速失败；
  2. 绿跑：comm 样例模块编译探针（code + 链接）→ 完整验证（构建/用例/覆盖率全绿）；
  3. 缺陷注入：把 CRC 初值改错 → 必须检出用例失败，且纯测试失败交归因（auto_decision=None）；
  4. 交付件：真实静态报告 + 真实执行结论 → 装配测评报告 → 导出 docx/xlsx 并确认可打开。

全部结论来自确定性工具输出。任一步不达标即退出码非 0，并打印失败明细。

用法：`.venv\\Scripts\\python.exe scripts\\e2e_verify_vm.py`
前置：先用 `scripts\\setup_verify_vm.ps1` 配好免密，.env 里 VERIFY_HOST/USER/KEY 就位。
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config
from core import c_static
from exporters.report_exporter import (assemble_report, export_report_docx,
                                       export_report_excel, report_markdown)
from pipeline.nodes import _slim_exec, _slim_static
from tests.fixtures import comm_sample as S
from verification import executor
from verification.runner import runner_from_config

# 用独立的项目号与工作区，绝不碰真实项目数据
PID = 9999
OUT_DIR = config.VERIFY_EVIDENCE_DIR / "e2e"

_FAILS: list[str] = []


def _check(cond: bool, label: str, detail: str = "") -> bool:
    if cond:
        print("  [OK]   " + label)
    else:
        print("  [FAIL] " + label + (("  → " + detail) if detail else ""))
        _FAILS.append(label)
    return bool(cond)


def _step(n: str, title: str) -> None:
    print("\n=== %s %s ===" % (n, title))


def step_probe(r) -> bool:
    _step("1/4", "探测验证机工具链")
    info = r.probe()
    _check(info.get("ok") is True, "工具链可用", info.get("error", ""))
    for k in ("uname", "cores", "gcc", "gcov", "make", "tar"):
        print("        %-6s %s" % (k, info.get(k)))
    _check(bool(info.get("gcc")) and bool(info.get("gcov")), "gcc 与 gcov 均在位")
    return info.get("ok") is True


def step_green(r) -> dict | None:
    _step("2/4", "绿跑：编译探针 + 完整验证（comm 样例模块）")
    p1 = executor.compile_probe(r, PID, S.CODE_FILES, None, stage="code")
    _check(p1["ok"] and not p1.get("skipped"), "code 阶段编译探针通过",
           p1.get("reason", ""))
    p2 = executor.compile_probe(r, PID, S.CODE_FILES, S.TEST_FILES, stage="test_impl")
    _check(p2["ok"] and not p2.get("skipped"), "test_impl 阶段编译+链接探针通过",
           p2.get("reason", ""))

    out = executor.run_verification(r, PID, "e2e-green", S.CODE_FILES, S.TEST_FILES,
                                    case_ids=S.CASE_IDS,
                                    branch_min=config.COVERAGE_BRANCH_MIN,
                                    line_min=config.COVERAGE_LINE_MIN)
    v = out["verdict"]
    t = out["tests"]
    cov = out["coverage"]
    _check(out["ok"] is True and out["decision"] == executor.DECISION_NEXT,
           "完整验证全绿（decision=next）", out.get("reason", ""))
    _check(v.get("build") == "ok", "构建判定 ok", str(v))
    _check(v.get("tests") == "ok" and t.get("all_pass") is True,
           "用例全过 %s/%s" % (t.get("passed"), t.get("total")), str(t.get("failed")))
    _check(t.get("missing") in (0, None, []), "无未执行用例", str(t.get("missing")))
    _check(v.get("coverage") == "ok" and cov.get("ok") is True,
           "覆盖率达标 分支 %.2f%%（门限 %.0f%%）"
           % (cov.get("totals", {}).get("branch_pct") or 0.0, config.COVERAGE_BRANCH_MIN))
    print("        行 %.2f%%（%s/%s）　分支 %.2f%%（%s/%s）" % (
        cov.get("totals", {}).get("line_pct") or 0.0,
        cov.get("totals", {}).get("lines_hit"), cov.get("totals", {}).get("lines_total"),
        cov.get("totals", {}).get("branch_pct") or 0.0,
        cov.get("totals", {}).get("branch_taken"), cov.get("totals", {}).get("branch_total")))
    return out if out["ok"] else None


def step_defect(r) -> None:
    _step("3/4", "缺陷注入：CRC 初值改错 → 必须检出并交归因")
    bad_c = S.COMM_C.replace("uint16_t crc = 0xFFFFu;", "uint16_t crc = 0x0000u;")
    if not _check(bad_c != S.COMM_C, "缺陷注入命中 CRC 初值"):
        return
    bad_code = {**S.CODE_FILES, "src/comm.c": bad_c}
    out = executor.run_verification(r, PID, "e2e-defect", bad_code, S.TEST_FILES,
                                    case_ids=S.CASE_IDS,
                                    branch_min=config.COVERAGE_BRANCH_MIN,
                                    line_min=config.COVERAGE_LINE_MIN)
    t = out["tests"]
    _check(out["ok"] is False, "缺陷被检出（ok=False）", out.get("reason", ""))
    _check(out["verdict"].get("build") == "ok", "构建仍 ok（缺陷是逻辑错，非编译错）")
    _check(out["verdict"].get("tests") == "fail" and (t.get("failed") or 0) > 0,
           "用例判定 fail，失败 %s 条" % t.get("failed"))
    # 纯测试失败（构建 ok / 覆盖 ok / 无编译诊断）→ 工具链判不死方向，交归因智能体
    _check(executor.auto_decision(out) is None,
           "auto_decision 交归因智能体（返回 None，不武断判方向）")
    brief = executor.failure_brief(out)
    _check(bool(brief.strip()) and "FAIL" in brief, "failure_brief 给出可用失败摘要")
    print("        失败摘要首行：" + (brief.splitlines() or [""])[0][:80])


def step_deliverables(r, green_out: dict | None) -> None:
    _step("4/4", "交付件：真实静态报告 + 执行结论 → docx/xlsx")
    srep = c_static.check_files(S.CODE_FILES, {"complexity_max": config.COMPLEXITY_MAX},
                                S.LLD_FUNCTIONS)
    _check(srep["ok"] and srep["required"] == 0, "静态检查零必查项违规",
           "required=%s" % srep["required"])

    out = green_out or executor.run_verification(
        r, PID, "e2e-rpt", S.CODE_FILES, S.TEST_FILES, case_ids=S.CASE_IDS,
        branch_min=config.COVERAGE_BRANCH_MIN, line_min=config.COVERAGE_LINE_MIN)
    arts = {
        "requirement": {"markdown": "# 需求规格\nFR-001..FR-006", "meta": {},
                        "version": 1, "status": "approved"},
        "hld": {"markdown": "# 概要设计", "meta": {}, "version": 1, "status": "approved"},
        "lld": {"markdown": "# 详细设计", "meta": {"functions": S.LLD_FUNCTIONS},
                "version": 1, "status": "approved"},
        "testcase": {"markdown": "# 测试用例设计",
                     "meta": {"cases": [{"id": c, "fr_ids": ["FR-001"]} for c in S.CASE_IDS]},
                     "version": 1, "status": "approved"},
        "code": {"markdown": "# 代码实现",
                 "meta": {"files": ["include/comm.h", "src/comm.c"]},
                 "version": 1, "status": "approved"},
        "test_impl": {"markdown": "# 测试实现",
                      "meta": {"cases": [{"id": c, "fn": "comm_pack", "fr_ids": ["FR-001"]}
                                         for c in S.CASE_IDS]},
                      "version": 1, "status": "approved"},
        "static": {"markdown": c_static.markdown_report(srep), "meta": _slim_static(srep),
                   "version": 1, "status": "approved"},
        "exec": {"markdown": executor.markdown_report(out), "meta": _slim_exec(out),
                 "version": 1, "status": "approved"},
    }
    versions = {k: v["version"] for k, v in arts.items()}
    data = assemble_report(arts, "通信协议样例模块（E2E 验收）", versions,
                           history={"static": [arts["static"]], "exec": [arts["exec"]]},
                           coverage_min=config.COVERAGE_BRANCH_MIN)
    _check((data.get("conclusion") or {}).get("pass") is True, "测评结论：通过")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    dx = str(OUT_DIR / "e2e_report.docx")
    xx = str(OUT_DIR / "e2e_report.xlsx")
    export_report_docx(data, dx)
    export_report_excel(data, xx)
    _check(Path(dx).exists() and Path(dx).stat().st_size > 0, "docx 已导出", dx)
    _check(Path(xx).exists() and Path(xx).stat().st_size > 0, "xlsx 已导出", xx)

    # 确认两类交付件真的能打开（不是写了个损坏文件）
    try:
        from openpyxl import load_workbook
        sheets = load_workbook(xx).sheetnames
        _check("证据清单" in sheets and len(sheets) >= 6, "xlsx 可打开且含证据清单等 6 表",
               str(sheets))
    except Exception as e:
        _check(False, "xlsx 可打开", repr(e))
    try:
        import docx
        d = docx.Document(dx)
        _check(len(d.paragraphs) > 0, "docx 可打开（%d 段 / %d 表）"
               % (len(d.paragraphs), len(d.tables)))
    except Exception as e:
        _check(False, "docx 可打开", repr(e))
    _check(len(report_markdown(data)) > 0, "report Markdown 可渲染")
    print("        交付件目录：" + str(OUT_DIR))


def main() -> int:
    print("航天嵌入式 AI 自动化代码验证系统 —— 判据层端到端验收")
    print("验证机：%s@%s:%s　工作区根：%s"
          % (config.VERIFY_USER, config.VERIFY_HOST, config.VERIFY_PORT, config.VERIFY_WORKDIR))
    if not config.verify_configured():
        print("\n[中止] 未配置验证机（VERIFY_HOST 为空）。先跑 scripts\\setup_verify_vm.ps1。")
        return 2
    r = runner_from_config()
    if r is None:
        print("\n[中止] 无法构造 runner。检查 .env 的 VERIFY_* 配置。")
        return 2

    try:
        if not step_probe(r):
            print("\n[中止] 验证机不可达或工具链缺失，后续步骤无意义。")
            return 1
        green = step_green(r)
        step_defect(r)
        step_deliverables(r, green)
    except Exception as e:
        import traceback
        print("\n[异常] " + type(e).__name__ + ": " + str(e))
        traceback.print_exc()
        return 1

    print("\n" + "=" * 56)
    if _FAILS:
        print("验收未通过：%d 项失败" % len(_FAILS))
        for f in _FAILS:
            print("  - " + f)
        return 1
    print("验收通过：判据层在真实验证机上全链路达标（绿跑 + 缺陷检出 + 交付件）。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
