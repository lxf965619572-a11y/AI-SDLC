"""生成星载通信协议样例模块的 PRD（docx），作为「生成 + 验证闭环」的试点输入。

为什么单独一份：`make_sample_prd.py` 的电商 PRD 是离线演示（mock）的输入，
mock 的详细设计/代码/测试都是按订单模块写死的，改掉它离线演示就散了。
这份 PRD 描述的是通信协议模块（成帧 / CRC / 解帧 / 确认窗口 / 重传 / 健壮性），
与 `tests/fixtures/comm_sample.py` 的基准代码同源：帧格式、窗口大小、错误码、
CRC 多项式都一一对应，因此既能喂真实模型跑端到端，也能拿 fixture 校准解析器。

用法：`.venv\\Scripts\\python.exe scripts\\make_comm_prd.py [输出路径]`
"""
import sys

from docx import Document

# 需求编号与 tests/fixtures/comm_sample.py 的 CASE_FR 对应关系一致：
# FR-001 成帧 / FR-002 CRC / FR-003 解帧 / FR-004 确认窗口 / FR-005 健壮性 / FR-006 重传
FRS = [
    ("FR-001", "成帧（发送封装）", "P0",
     "给定不超过 32 字节的载荷，输出一帧完整报文，并把该帧登记到重传窗口"),
    ("FR-002", "CRC-16 校验计算", "P0",
     "对 SEQ+LEN+PAYLOAD 计算 CRC-16/CCITT-FALSE，标准校验串 123456789 得 0x29B1"),
    ("FR-003", "解帧（接收解析）", "P0",
     "校验帧头、长度、CRC 与接收序号，全部通过才交付载荷并推进期望序号"),
    ("FR-004", "确认与重传窗口", "P0",
     "窗口容量 4 帧；收到确认后移除对应帧并压缩窗口；窗口满时拒绝新帧"),
    ("FR-005", "入参防护与健壮性", "P0",
     "全部对外接口对空指针、非法长度、缓冲区不足返回明确错误码，不崩溃、不改动状态"),
    ("FR-006", "重传取帧", "P1",
     "取出窗口中最旧的未确认帧供重传，窗口为空时返回明确的空错误并累计重传次数"),
]

FRAME_FIELDS = [
    ("HEAD", "1", "帧头，固定 0xA5"),
    ("SEQ", "1", "发送序号，0..255 循环递增，回绕后从 0 重新开始"),
    ("LEN", "1", "载荷字节数，取值 0..32"),
    ("PAYLOAD", "LEN", "载荷数据，最大 32 字节"),
    ("CRC16", "2", "大端序，覆盖 SEQ、LEN 与 PAYLOAD，不含 HEAD"),
]

ERROR_CODES = [
    ("0", "成功"),
    ("-1", "参数非法：空指针、帧头不匹配等"),
    ("-2", "长度非法：载荷超长、报文长度不足、输出缓冲区容量不够"),
    ("-3", "CRC 校验失败"),
    ("-4", "序号错误：与期望接收序号不一致"),
    ("-5", "重传窗口已满"),
    ("-6", "重传窗口为空"),
]


def _table(doc, header, rows, style="Table Grid"):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(header))
    tbl.style = style
    for j, h in enumerate(header):
        tbl.rows[0].cells[j].text = h
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            tbl.rows[i].cells[j].text = str(v)
    return tbl


def build() -> Document:
    doc = Document()
    doc.add_heading("星载通信协议模块 PRD（验证闭环试点）", level=1)

    doc.add_heading("1. 项目背景", level=2)
    doc.add_paragraph(
        "某型号星载数传分系统需要一套遥测帧收发模块，负责把应用数据成帧下发、"
        "对上行帧做校验与解帧，并在链路误码时按窗口重传。模块运行在裸机环境，"
        "无操作系统、无堆分配，代码必须可静态分析、可单元验证。")
    doc.add_paragraph(
        "本模块同时是「航天嵌入式 AI 自动化代码验证系统」的试点目标：从本 PRD 出发，"
        "依次产出需求规格、概要设计、详细设计、测试用例设计、C 源码与可执行测试，"
        "再由静态检查、远端编译执行与覆盖率统计给出确定性结论，最后汇编成交付件。")

    doc.add_heading("2. 帧格式", level=2)
    doc.add_paragraph(
        "帧总长 = LEN + 5 字节，最大 37 字节。字段定义如下，多字节字段除 CRC 外均单字节：")
    _table(doc, ["字段", "长度（字节）", "说明"], FRAME_FIELDS)
    doc.add_paragraph(
        "CRC 采用 CRC-16/CCITT-FALSE：多项式 0x1021，初值 0xFFFF，输入与输出均不反转，"
        "无最终异或。计算范围是 SEQ、LEN 与 PAYLOAD 三段连续字节，结果按大端序附在帧尾。")

    doc.add_heading("3. 功能需求", level=2)
    _table(doc, ["编号", "需求名称", "优先级", "验收判据"], FRS)

    doc.add_heading("3.1 FR-001 成帧（发送封装）", level=3)
    for t in [
        "输入：待发载荷指针与长度、输出缓冲区指针与其容量、通信上下文。",
        "输出缓冲区写入 HEAD、SEQ、LEN、PAYLOAD、CRC16（大端），返回帧总长 LEN + 5。",
        "SEQ 取上下文中的下一个发送序号，成帧成功后发送序号加 1，超过 255 回绕到 0。",
        "成帧成功的帧必须同时登记到重传窗口，未确认前不得被丢弃。",
        "载荷长度为 0 是合法输入，应产出一条只含帧头帧尾的空载荷帧。",
        "载荷长度超过 32 字节、或输出缓冲区容量小于 LEN + 5 时，拒绝成帧并返回长度错误，"
        "不得向输出缓冲区写入任何字节，也不得改动发送序号与窗口。",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("3.2 FR-002 CRC-16 校验计算", level=3)
    for t in [
        "算法：CRC-16/CCITT-FALSE，多项式 0x1021，初值 0xFFFF，不反转，无最终异或。",
        "标准校验串 \"123456789\"（9 字节 ASCII）的计算结果必须是 0x29B1。",
        "空数据或长度为 0 时返回初值 0xFFFF，不得崩溃。",
        "实现不得依赖查表以外的运行期初始化，禁止动态分配查找表。",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("3.3 FR-003 解帧（接收解析）", level=3)
    for t in [
        "输入：原始字节流指针与其长度、输出帧结构指针、通信上下文。",
        "判据顺序固定：先查空指针与最小长度（不足 5 字节直接判长度错误），"
        "再查帧头是否为 0xA5，再查报文长度是否够 LEN + 5，然后校验 CRC，最后校验序号。",
        "CRC 不匹配返回 -3；CRC 正确但 SEQ 与期望接收序号不一致返回 -4。",
        "只有全部判据通过才写出载荷、返回成功，并把期望接收序号加 1（超过 255 回绕到 0）。",
        "任何一步失败都不得改动上下文中的期望接收序号，也不得写出输出帧。",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("3.4 FR-004 确认与重传窗口", level=3)
    for t in [
        "窗口容量固定 4 帧，是编译期常量，不得在运行期改写。",
        "已发送未确认的帧按发送顺序存放在窗口里，窗口满时新的成帧请求返回 -5。",
        "收到确认序号后，从窗口中移除该序号对应的帧，其后各帧前移一位，窗口计数减 1。",
        "确认一个不在窗口内的序号是无害操作：窗口内容与计数保持不变，不得崩溃。",
        "同一序号被重复确认时，第二次同样是无害操作。",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("3.5 FR-005 入参防护与健壮性", level=3)
    for t in [
        "全部对外接口必须做入参防护：空指针、负长度、超长长度、缓冲区容量不足、"
        "非法帧头，一律返回明确错误码，不得解引用空指针、不得越界读写。",
        "失败路径不得留下半成品状态：上下文中的序号、窗口计数、已登记帧内容保持调用前的值。",
        "接口不得依赖调用顺序之外的隐式前提；上下文未初始化时的行为由调用方负责，"
        "但模块自身不得因为字段取值异常而崩溃。",
        "模块不得有文件作用域可写变量，全部运行期状态收敛在调用方传入的上下文结构体中。",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("3.6 FR-006 重传取帧", level=3)
    for t in [
        "取帧接口返回窗口中最旧的未确认帧（即最早成帧且尚未被确认的那一帧）的内容，"
        "并返回成功；取出后该帧仍留在窗口中，等待确认。",
        "窗口为空时返回 -6，不得写出输出帧。",
        "每次成功取帧，上下文中的累计重传次数加 1，计数到 255 后回绕。",
        "重传取帧不得改变发送序号，也不得改变窗口中帧的顺序。",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("4. 接口与实现约束", level=2)
    for t in [
        "语言与编译：C99，以 `gcc -std=c99 -Wall -Wextra` 零告警为门槛。",
        "对外接口以 `comm_` 前缀命名，原型与数据结构全部声明在 `include/comm.h`，"
        "实现放在 `src/comm.c`，内部辅助函数必须声明为 static。",
        "通信上下文 `CommCtx` 由调用方持有并显式传入，模块内不得出现全局变量或 static 局部变量。",
        "返回值约定：成功返回 0（或帧长度等明确非负值），失败返回负错误码。",
        "测试只包含头文件即可编写（黑盒），不得依赖实现内部符号。",
    ]:
        doc.add_paragraph(t, style="List Bullet")
    doc.add_paragraph("错误码取值：")
    _table(doc, ["返回码", "含义"], ERROR_CODES)

    doc.add_heading("5. 非功能需求", level=2)
    for t in [
        "禁止动态内存分配、禁止递归、禁止函数指针、禁止变长数组；所有缓冲区尺寸是编译期常量。",
        "单函数圈复杂度不超过 10，超出必须拆分。",
        "单帧处理不得阻塞、不得调用操作系统服务，最坏栈深度可静态确定。",
        "分支覆盖率不低于 80%，且不允许存在完全未被测试执行的对外函数。",
        "每条功能需求至少有一条可执行测试用例，判据取自本文档而非实现。",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("6. 验收方式", level=2)
    doc.add_paragraph(
        "以自动验证闭环的结论为准：静态检查无必查项违规（或违规已走偏差单并经人工批准）、"
        "远端编译链接通过、全部测试用例执行通过、覆盖率达标。任一环节不达标即判不通过，"
        "并产出问题报告；结论与证据（原始日志、覆盖率报告、输入 sha256）一并归档到需求追溯矩阵。")
    return doc


def main() -> None:
    out = sys.argv[1] if len(sys.argv) > 1 else "comm_prd.docx"
    build().save(out)
    print("saved", out)
    for fid, name, prio, acc in FRS:
        print(f"  {fid} [{prio}] {name}：{acc}")


if __name__ == "__main__":
    main()
