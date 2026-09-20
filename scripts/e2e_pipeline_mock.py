"""全链路编排验收：mock 模型 + 真实验证机，一条命令跑完 parse→…→report。

为什么还要这一个脚本：`scripts/e2e_verify_vm.py` 验的是「判据层」——静态检查、编译、
用例、覆盖率这些确定性结论对不对；但系统能不能真按 V 模型串起来、6 道评审门是否照常被
interrupt 拦住、工具节点的结论能否回灌追溯矩阵并装配成交付件，只有把整张图跑一遍才知道。
真实模型跑一遍要人守 6 道门、要烧 token，不适合当回归；这里用 LLM_MOCK=1 把模型换成
core/mock_c.py 的固定产物，判据层仍然打真实验证机，于是可重复、零 token、无需 Web UI。

三个场景（对应方案「验收」里的两条硬要求：全链路绿 + 注入缺陷验闭环）：
  green      正常链路：6 道评审门 + 3 个工具节点，一路全绿到交付件；
  defect     注入一处代码缺陷（金卡折扣算成 9 折，TC-006 的判据对不上）：
             exec 失败 → 归因智能体定责为代码缺陷 → 回代码阶段重生 → 再过一次代码
             评审门 → 静态检查 → exec 转绿；这一轮失败在问题报告单里留痕并标注已闭环。
             只改坏一处常量表达式：编译过、静态过、覆盖率不变，纯粹是判据对不上，
             正好落在「确定性判据不足以定责、必须读代码与用例」的归因智能体分支上。
  overlimit  缺陷修不掉（每轮都生成同一份坏代码）：数满 MAX_FIX_ROUNDS 轮自动整改后
             出问题报告单，停在 gate_exec 由人工裁决；受理后仍然装配交付件，但测评
             结论必须是「不通过」——失败绝不静默放行。

数据隔离：临时 SQLite（data/e2e_pipeline/），绝不碰 data/app.sqlite 里的真实项目。

用法：`.venv\\Scripts\\python.exe scripts\\e2e_pipeline_mock.py [green|defect|overlimit]`
前置：scripts\\setup_verify_vm.ps1 已配好免密；仓库根有 comm_prd.docx（脚本自动生成时可
先跑 scripts\\make_comm_prd.py）。
"""
from __future__ import annotations

import os
import shutil
import sys
from pathlib import Path

# 必须在导入任何会读它的模块之前置位：离线演示产物由 core.mock_llm 按 role 返回
os.environ["LLM_MOCK"] = "1"

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config

# ---- 数据隔离：换库要在 db.models 建 engine 之前 ----
ROOT = config.DATA_DIR / "e2e_pipeline"
ROOT.mkdir(parents=True, exist_ok=True)
(ROOT / "uploads").mkdir(exist_ok=True)
(ROOT / "verify").mkdir(exist_ok=True)
for _old in ("app.sqlite", "checkpoints.sqlite"):
    _p = ROOT / _old
    if _p.exists():
        _p.unlink()
config.APP_DB = ROOT / "app.sqlite"
config.CHECKPOINT_DB = ROOT / "checkpoints.sqlite"
config.UPLOAD_DIR = ROOT / "uploads"
config.VERIFY_EVIDENCE_DIR = ROOT / "verify"

import db.models as M                                              # noqa: E402
from core import mock_c, trace as T                                # noqa: E402
from exporters.report_exporter import (export_report_docx,          # noqa: E402
                                       export_report_excel)
from exporters.trace_exporter import export_trace_excel             # noqa: E402
from pipeline import nodes                                          # noqa: E402
from services.report_service import build_report_data               # noqa: E402
from services.trace_service import build_project_matrix             # noqa: E402
from verification import executor                                   # noqa: E402

SCENARIOS = ("green", "defect", "overlimit")
SCENARIO_TITLES = {
    "green": "正常链路：一路全绿到交付件",
    "defect": "注入代码缺陷：失败 → 归因 → 重生 → 过门 → 转绿",
    "overlimit": "缺陷修不掉：数满自动整改轮数 → 问题报告单 → 人工裁决",
}

_FAILS: list[str] = []
_SCEN = "green"


def _check(cond: bool, label: str, detail: str = "") -> bool:
    if cond:
        print("  [OK]   " + label)
    else:
        print("  [FAIL] " + label + (("  → " + detail) if detail else ""))
        _FAILS.append(f"{_SCEN}: {label}")
    return bool(cond)


def _step(n: str, title: str) -> None:
    print("\n--- %s %s" % (n, title))


# ---------------- 缺陷注入 ----------------
# 只改折扣计算这一行：95 折变 90 折。编译照过、静态照过、分支结构不变（覆盖率不动），
# 唯一后果是 TC-006「金卡会员 95 折」的判据对不上——正是需要归因智能体读代码定责的那类
# 缺陷，也是航天软件里最典型的「实现与设计常量不一致」。
DEFECT_OLD = "paid = (amount_cent * ORDER_DISCOUNT_NUM) / ORDER_DISCOUNT_DEN;"
DEFECT_NEW = "paid = (amount_cent * (ORDER_DISCOUNT_NUM - 5)) / ORDER_DISCOUNT_DEN;"


def defective_source() -> str:
    if DEFECT_OLD not in mock_c.ORDER_C:
        raise SystemExit("缺陷注入点失效：mock 源码里找不到折扣计算行，请同步 DEFECT_OLD")
    return mock_c.ORDER_C.replace(DEFECT_OLD, DEFECT_NEW)


def install_defect(scenario: str):
    """把「代码阶段第 1 次生成」换成缺陷版；overlimit 场景则每次都换。

    注入点选在 core.llm_client.mock_complete（离线模式下所有阶段都从这里出），
    只在 role == "code" 时临时替换 core.mock_c.CODE_FILES——归因、测试实现、
    交付件读到的仍是库里那份真实产物，不受影响。返回还原函数。"""
    if scenario == "green":
        return lambda: None
    import core.llm_client as LC

    bad_files = {"include/order.h": mock_c.ORDER_H, "src/order.c": defective_source()}
    always = scenario == "overlimit"
    orig = LC.mock_complete
    calls = {"code": 0}

    def patched(messages, role):
        if role != "code":
            return orig(messages, role)
        calls["code"] += 1
        if not always and calls["code"] > 1:
            return orig(messages, role)
        saved = mock_c.CODE_FILES
        mock_c.CODE_FILES = bad_files
        try:
            return orig(messages, role)
        finally:
            mock_c.CODE_FILES = saved

    LC.mock_complete = patched

    def undo():
        LC.mock_complete = orig

    return undo


# ---------------- 驱动 ----------------
def make_project(scenario: str) -> int:
    """建一个干净项目并把样例 PRD 挂上去（parse 阶段要读 Document 行）。"""
    prd = config.BASE_DIR / "comm_prd.docx"
    if not prd.exists():
        raise SystemExit("缺少样例 PRD：%s（先跑 scripts\\make_comm_prd.py）" % prd)
    M.init_db()
    with M.SessionLocal() as s:
        proj = M.Project(name="E2E 全链路验收 · %s（mock 模型 + 真实验证机）"
                         % SCENARIO_TITLES[scenario].split("：")[0], status="created")
        s.add(proj)
        s.commit()
        pid = proj.id
        dst = config.UPLOAD_DIR / prd.name
        shutil.copyfile(prd, dst)
        s.add(M.Document(project_id=pid, filename=prd.name,
                         stored_path=str(dst), file_type="docx"))
        s.commit()
    return pid


def run_graph(pid: int) -> tuple[list[str], bool, dict]:
    """同步驱动整张图：遇 interrupt（评审门）就自动通过，直到 END。

    返回 (被拦下的评审门阶段序列, 是否正常收尾, 终态)。停在非评审门节点或回环失控都算
    失败——前者意味着某个节点抛错后图卡住，后者意味着失败闭环的轮数上限没起作用。"""
    from langgraph.checkpoint.sqlite import SqliteSaver
    from langgraph.types import Command

    from pipeline.graph import build_graph

    gates: list[str] = []
    cfg = {"configurable": {"thread_id": "e2e_%d" % pid}}
    with SqliteSaver.from_conn_string(str(config.CHECKPOINT_DB)) as cp:
        graph = build_graph(cp)
        graph.invoke({"project_id": pid}, config=cfg)
        for _ in range(40):
            snap = graph.get_state(cfg)
            if not snap.next:
                break
            ints = [i for t in snap.tasks for i in (getattr(t, "interrupts", None) or [])]
            if not ints:
                _check(False, "图停在评审门以外的节点", str(snap.next))
                return gates, False, snap.values or {}
            payload = ints[0].value or {}
            gates.append(payload.get("stage") or "?")
            print("        评审门 %-10s v%s → 自动通过"
                  % (payload.get("stage"), payload.get("version")))
            graph.invoke(Command(resume={"approved": True,
                                         "comments": "E2E 自动通过"}), config=cfg)
        else:
            _check(False, "评审门超过 40 次，失败闭环疑似失控")
            return gates, False, {}
        snap = graph.get_state(cfg)
        return gates, not snap.next, snap.values or {}


# ---------------- 库里取数 ----------------
def stage_versions(pid: int, stage: str) -> list[dict]:
    with M.SessionLocal() as s:
        rows = (s.query(M.StageArtifact).filter_by(project_id=pid, stage=stage)
                .order_by(M.StageArtifact.version.asc()).all())
        return [{"version": r.version, "status": r.status, "markdown": r.markdown or "",
                 "meta": r.meta_json or {}} for r in rows]


def latest_meta(pid: int, stage: str) -> dict:
    with M.SessionLocal() as s:
        a = (s.query(M.StageArtifact).filter_by(project_id=pid, stage=stage)
             .order_by(M.StageArtifact.version.desc()).first())
        return (a.meta_json or {}) if a else {}


def log_lines(pid: int, keyword: str) -> list[str]:
    with M.SessionLocal() as s:
        rows = s.query(M.PipelineLog).filter_by(project_id=pid).all()
        return [r.message for r in rows if keyword in (r.message or "")]


def failed_cases(meta: dict) -> list[str]:
    return [r.get("id") for r in ((meta.get("tests") or {}).get("results") or [])
            if r.get("status") != "pass"]


# ---------------- 各场景判据 ----------------
def check_graph_shape(scenario: str, pid: int, gates: list[str], finished: bool,
                      values: dict) -> None:
    from services import pipeline_service as PS

    expect_gates = {
        "green": list(nodes.STAGES),
        # 自动整改的每一版代码都要重新过评审门（AI 改代码的权限边界）
        "defect": list(nodes.STAGES) + ["code"],
        "overlimit": list(nodes.STAGES) + ["code"] * config.MAX_FIX_ROUNDS + ["exec"],
    }[scenario]
    _check(finished, "图跑到 END（无残留待执行节点）")
    _check(gates == expect_gates, "评审门序列 = %s" % ",".join(expect_gates), str(gates))
    with M.SessionLocal() as s:
        proj = s.get(M.Project, pid)
        status = proj.status if proj else "?"
        stage = proj.current_stage if proj else "?"
        reviews = s.query(M.Review).filter_by(project_id=pid).count()
        arts = {a.stage: a for a in
                s.query(M.StageArtifact).filter_by(project_id=pid).all()}
    # 置 completed 是服务层 _run_graph 收尾时做的（本脚本同步驱动图，不经过那一行），
    # 所以这里查图自己负责的事实：走到最后一个阶段、既没停在评审门也没有可续跑节点。
    _check(stage == "report" and status in ("running", "completed"),
           "流水线推进到 report 阶段（状态 %s）" % status, "%s/%s" % (status, stage))
    _check(PS.detect_interrupted_stage(pid) is None, "没有滞留在评审门")
    _check(PS.resumable_stage(pid) is None, "没有待续跑节点（确实终止，不是卡住）")
    _check(reviews >= len(expect_gates), "评审记录留痕 %d 条" % reviews, str(reviews))
    missing = [st for st in nodes.DISPLAY_STAGES if st not in arts]
    _check(not missing, "十个阶段产物齐备：%s" % ",".join(nodes.DISPLAY_STAGES),
           "缺 " + ",".join(missing))
    bad_status = sorted({a.stage for a in arts.values() if a.status != "approved"})
    _check(not bad_status, "产物最终状态均为 approved（超限项经人工裁决后受理）",
           str(bad_status))
    fr = values.get("fix_rounds") or {}
    if scenario == "defect":
        # 计数能活着穿过代码评审门，超限判据才成立（门的额度语义见 nodes.gate_fix_rounds）
        _check(fr.get("exec") == 1, "自动整改轮数记账 = 1（未被评审门清零）", str(fr))


def check_tools(scenario: str, pid: int) -> None:
    sm = latest_meta(pid, "static")
    _check(sm.get("ok") is True and sm.get("required") == 0,
           "静态检查零必查项违规（总违规 %s）" % sm.get("total"), str(sm.get("required")))
    execs = stage_versions(pid, "exec")
    codes = stage_versions(pid, "code")
    expect_runs = {"green": 1, "defect": 2, "overlimit": config.MAX_FIX_ROUNDS + 1}[scenario]
    _check(len(execs) == expect_runs, "执行验证跑了 %d 轮" % expect_runs,
           "实际 %d 轮" % len(execs))

    if scenario != "green":
        first = execs[0]["meta"]
        _check(first.get("ok") is False, "第 1 轮判据为未通过（缺陷被工具链抓到）",
               str(first.get("verdict")))
        _check((first.get("verdict") or {}).get("build") == "ok",
               "缺陷不影响编译：构建仍 ok，失败落在用例判据上")
        _check("TC-006" in failed_cases(first), "失败用例点名 TC-006（金卡折扣）",
               str(failed_cases(first)))
        att = first.get("attribution") or {}
        _check(att.get("decision") == executor.DECISION_FIX_CODE
               and att.get("source") == "agent",
               "归因智能体定责：被测代码缺陷（source=%s）" % att.get("source"), str(att))
        _check(bool(log_lines(pid, "归因（智能体）")), "归因过程在流水线日志里留痕")
        expect_code_versions = expect_runs if scenario == "overlimit" else 2
        _check(len(codes) == expect_code_versions,
               "代码重生 %d 版" % expect_code_versions, "实际 %d 版" % len(codes))
        if scenario == "overlimit":
            _check(all(DEFECT_NEW in c["markdown"] for c in codes),
                   "缺陷标记每轮都在（一直修不掉）",
                   str([DEFECT_NEW in c["markdown"] for c in codes]))
        else:
            _check(DEFECT_NEW in codes[0]["markdown"]
                   and DEFECT_NEW not in codes[-1]["markdown"],
                   "缺陷标记出现在第 1 版代码里、重生后消失")

    last = execs[-1]["meta"]
    v, t, cov = last.get("verdict") or {}, last.get("tests") or {}, last.get("coverage") or {}
    if scenario == "overlimit":
        _check(last.get("ok") is False, "最后一轮仍未通过（没有偷偷放行）")
        _check("问题报告单" in execs[-1]["markdown"]
               and f"自动修复轮数：{config.MAX_FIX_ROUNDS}" in execs[-1]["markdown"],
               "执行报告里落了问题报告单（含轮数上限）")
        _check(bool(log_lines(pid, "出问题报告单转人工裁决")), "转人工裁决写进日志")
    else:
        _check(last.get("ok") is True and last.get("decision") == executor.DECISION_NEXT,
               "执行验证全绿（decision=%s）" % last.get("decision"), last.get("reason", ""))
        _check(v.get("build") == "ok" and v.get("tests") == "ok"
               and v.get("coverage") == "ok", "三项判定均 ok", str(v))
        _check(t.get("all_pass") is True,
               "用例 %s/%s 全过" % (t.get("passed"), t.get("total")), str(t.get("failed")))
        _check(cov.get("ok") is True,
               "覆盖率达标 分支 %.2f%%（门限 %.0f%%）"
               % ((cov.get("totals") or {}).get("branch_pct") or 0.0,
                  config.COVERAGE_BRANCH_MIN), str(cov.get("totals")))
    env = last.get("env") or {}
    _check(bool(env.get("gcc")) and bool(env.get("gcov")),
           "证据里记了验证机工具链版本：%s" % (env.get("gcc") or "")[:40])
    manifest = last.get("manifest") or {}
    _check(len(manifest) >= 3, "同步输入指纹 %d 项（同步前在本机算出 sha256）"
           % len(manifest), str(sorted(manifest)))
    arch = last.get("evidence")
    _check(bool(arch) and Path(str(arch)).exists(), "原始输出已归档：%s" % (arch or "-"),
           str(arch))
    if scenario != "green":
        tags = {Path(str(e["meta"].get("evidence") or "")).parent.name for e in execs}
        _check(len(tags) == len(execs), "每轮证据各自成目录（不覆盖历史）：%s"
               % ",".join(sorted(tags)), str(tags))


def check_trace(scenario: str, pid: int) -> None:
    mx = build_project_matrix(pid)
    rows = mx.get("rows") or []
    summary = mx.get("summary") or {}
    _check(len(rows) > 0, "矩阵 %d 行需求" % len(rows))
    empt = {c: [r["id"] for r in rows if r.get(c) is None]
            for c in ("code_units", "static_violations", "exec_result", "branch_coverage")}
    for c, ids in empt.items():
        _check(not ids, "列 %-18s 全部有值" % c, "空值行 " + str(ids))
    st = {r["id"]: T.row_status(r, summary) for r in rows}
    if scenario == "overlimit":
        bad = [k for k, v in st.items() if v == T.ST_EXEC_FAIL]
        _check(bool(bad), "失败如实反映到矩阵：至少一行状态 = 执行失败", str(st))
        _check(all(v in (T.ST_OK, T.ST_EXEC_FAIL) for v in st.values()),
               "其余行未被牵连（只挂 TC-006 那条需求）", str(st))
    else:
        _check(all(r.get("exec_result") == "all" for r in rows), "每行执行结果 = 全部通过",
               str({r["id"]: r.get("exec_result") for r in rows}))
        _check(all(v == T.ST_OK for v in st.values()),
               "每行状态 = 贯通（未被执行/覆盖维度降级）", str(st))
    xp = str(ROOT / f"{scenario}_trace_matrix.xlsx")
    export_trace_excel(mx, xp, "E2E 全链路验收 · " + scenario)
    _check(Path(xp).exists() and Path(xp).stat().st_size > 0, "追溯矩阵 xlsx 已导出")


def check_deliverables(scenario: str, pid: int) -> None:
    data = build_report_data(pid)
    concl = data.get("conclusion") or {}
    probs = data.get("problems") or []
    if scenario == "green":
        _check(concl.get("pass") is True, "测评结论：通过", str(concl.get("text")))
        _check(not probs, "无问题报告单 %d 项" % len(probs), str(probs))
    elif scenario == "defect":
        _check(concl.get("pass") is True, "修好后测评结论：通过", str(concl.get("text")))
        _check(len(probs) == 1 and probs[0]["stage"] == "exec",
               "失败那一轮登记为问题报告单 PR-001", str(probs))
        _check(probs and probs[0].get("resolved") is True
               and probs[0].get("decision") == executor.DECISION_FIX_CODE,
               "问题报告单标注「已闭环」并记住责任方判定",
               str(probs[0] if probs else None))
    else:
        _check(concl.get("pass") is False, "测评结论：不通过（失败不得静默放行）",
               str(concl.get("text")))
        _check(len(probs) == config.MAX_FIX_ROUNDS + 1,
               "每一轮失败都登记：%d 条问题报告单" % (config.MAX_FIX_ROUNDS + 1),
               str([p["id"] for p in probs]))
        _check(probs and probs[-1].get("resolved") is False,
               "最后一条标注「未闭环，转人工裁决」", str(probs[-1] if probs else None))
        _check(not (data.get("deviations") or []), "本次无静态偏差单（缺陷不在静态维度）")
    _check(bool((data.get("evidence") or [])), "证据清单非空 %d 项"
           % len(data.get("evidence") or []))
    dx = str(ROOT / f"{scenario}_report.docx")
    xx = str(ROOT / f"{scenario}_report.xlsx")
    export_report_docx(data, dx)
    export_report_excel(data, xx)
    try:
        from openpyxl import load_workbook
        sheets = load_workbook(xx).sheetnames
        _check(len(sheets) >= 6 and "证据清单" in sheets, "测评报告 xlsx 可打开（%d 表）"
               % len(sheets), str(sheets))
    except Exception as e:
        _check(False, "测评报告 xlsx 可打开", repr(e))
    try:
        import docx
        d = docx.Document(dx)
        _check(len(d.paragraphs) > 0, "测评报告 docx 可打开（%d 段 / %d 表）"
               % (len(d.paragraphs), len(d.tables)))
    except Exception as e:
        _check(False, "测评报告 docx 可打开", repr(e))
    # report 节点自己也落了一份产物，两者结论必须一致（装配逻辑只有一处）
    rm = latest_meta(pid, "report")
    _check((rm.get("conclusion") or {}).get("pass") is concl.get("pass"),
           "report 节点落库产物结论与导出一致（pass=%s）" % concl.get("pass"),
           str(rm.get("conclusion")))


def run_scenario(scenario: str) -> None:
    global _SCEN
    _SCEN = scenario
    print("\n" + "=" * 68)
    print("场景 %s —— %s" % (scenario, SCENARIO_TITLES[scenario]))
    print("=" * 68)
    pid = make_project(scenario)
    print("项目号：%d（临时库，不影响真实数据）" % pid)
    undo = install_defect(scenario)
    _step("1/4", "驱动整张图：评审门 + 工具节点 + 失败闭环")
    try:
        gates, finished, values = run_graph(pid)
    except Exception as e:
        import traceback
        traceback.print_exc()
        _check(False, "图执行未抛异常", "%s: %s" % (type(e).__name__, e))
        return
    finally:
        undo()
    check_graph_shape(scenario, pid, gates, finished, values)
    _step("2/4", "工具节点结论：静态检查 + 真实机执行")
    check_tools(scenario, pid)
    _step("3/4", "追溯矩阵：四个新列由真实产物填满")
    check_trace(scenario, pid)
    _step("4/4", "交付件：从库里真实产物装配并导出")
    check_deliverables(scenario, pid)
    print("        输出目录：%s（%s_*）" % (ROOT, scenario))


def main(argv: list[str]) -> int:
    chosen = [a for a in argv if not a.startswith("-")]
    unknown = [a for a in chosen if a not in SCENARIOS]
    if unknown:
        print("未知场景：%s（可选 %s）" % (",".join(unknown), " / ".join(SCENARIOS)))
        return 2
    chosen = chosen or list(SCENARIOS)

    print("航天嵌入式 AI 自动化代码验证系统 —— 全链路编排验收（mock 模型 + 真实验证机）")
    print("验证机：%s@%s:%s　临时库：%s"
          % (config.VERIFY_USER, config.VERIFY_HOST, config.VERIFY_PORT, ROOT))
    print("场景：%s　自动整改轮数上限：%d" % ("、".join(chosen), config.MAX_FIX_ROUNDS))
    if not config.verify_configured():
        print("\n[中止] 未配置验证机（VERIFY_HOST 为空）。先跑 scripts\\setup_verify_vm.ps1。")
        return 2

    for scenario in chosen:
        run_scenario(scenario)

    print("\n" + "=" * 68)
    if _FAILS:
        print("验收未通过：%d 项失败" % len(_FAILS))
        for f in _FAILS:
            print("  - " + f)
        return 1
    print("验收通过：%s 场景全部跑通（评审门 + 工具节点 + 失败闭环 + 追溯 + 交付件）。"
          % "、".join(chosen))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
