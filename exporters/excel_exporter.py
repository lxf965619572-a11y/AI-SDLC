"""测试用例导出 Excel（openpyxl）。"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["编号", "模块", "用例标题", "类型", "优先级", "前置条件", "测试步骤", "预期结果"]
WIDTHS = [10, 12, 30, 8, 8, 24, 50, 40]


def export_excel(testcases: list[dict], path: str, sheet_title: str = "测试用例"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(color="FFFFFF", bold=True)
    for col, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = w

    for row, c in enumerate(testcases, 2):
        steps = "\n".join(f"{i}. {s}" for i, s in enumerate(c.get("steps", []), 1))
        values = [c.get("id", ""), c.get("module", ""), c.get("title", ""),
                  c.get("type", ""), c.get("priority", ""), c.get("preconditions", ""),
                  steps, c.get("expected", "")]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=(col >= 6))

    ws.freeze_panes = "A2"
    wb.save(path)
    return path
