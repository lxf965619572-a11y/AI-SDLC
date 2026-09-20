"""需求追溯矩阵导出 Excel（openpyxl）。

矩阵的计算全在 core.trace、取数在 services.trace_service，这里只负责排版：
Sheet1 按需求逐行展开「素材 → 需求 → 概设 → 详设 → 用例」，Sheet2 是覆盖概览。
两张表都是交付件口径，可直接归档或送审。
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core import trace

HEADERS = ["需求编号", "优先级", "需求描述", "素材来源",
           "概要设计", "详细设计", "测试用例", "链路状态"]
WIDTHS = [11, 8, 46, 30, 24, 30, 26, 30]

# 状态配色：待补全最抢眼；未记录最弱（是数据缺失，不是缺陷）
STATUS_STYLE = {
    trace.ST_GAP: ("FFF2CC", "9C5700"),
    trace.ST_UNRECORDED: ("EDEDED", "6B6B6B"),
    trace.ST_PENDING: ("DEEBF7", "1F4E79"),
    trace.ST_OK: ("E2EFDA", "375623"),
}
STATUS_ORDER = (trace.ST_GAP, trace.ST_PENDING, trace.ST_UNRECORDED, trace.ST_OK)

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(color="FFFFFF", bold=True)

STAGE_LABELS = {"requirement": "需求规格", "hld": "概要设计",
                "lld": "详细设计", "testcase": "测试用例"}


def _sources_text(row: dict) -> str:
    """素材编号 + 名称，一行一个；没有名称映射时退回纯编号。"""
    ids = row.get("sources") or []
    if not ids:
        return ""
    names = row.get("source_names") or []
    lines = []
    for i, sid in enumerate(ids):
        name = names[i] if i < len(names) else ""
        lines.append(f"{sid} {name}".strip() if name and name != sid else str(sid))
    return "\n".join(lines)


def _header_row(ws, headers, widths):
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = w


def _matrix_sheet(wb, rows):
    ws = wb.active
    ws.title = "追溯矩阵"
    _header_row(ws, HEADERS, WIDTHS)

    for r, row in enumerate(rows, 2):
        values = [row.get("id", ""), row.get("priority", ""), row.get("desc", ""),
                  _sources_text(row),
                  "\n".join(row.get("hld_modules") or []),
                  "\n".join(row.get("lld_functions") or []),
                  "\n".join(row.get("testcases") or []),
                  row.get("status_text", "")]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=col >= 3)
        style = STATUS_STYLE.get(row.get("status"))
        if style:
            cell = ws.cell(row=r, column=len(HEADERS))
            cell.fill = PatternFill("solid", fgColor=style[0])
            cell.font = Font(color=style[1], bold=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(len(rows) + 1, 1)}"
    return ws


def _overview_sheet(wb, data, project_name):
    rows = data.get("rows") or []
    summary = data.get("summary") or {}
    versions = data.get("versions") or {}
    gaps = summary.get("uncovered_p0") or {}
    covered = summary.get("covered") or {}
    p0 = summary.get("p0_total") or 0

    status_count = {}
    for row in rows:
        key = row.get("status") or trace.ST_OK
        status_count[key] = status_count.get(key, 0) + 1

    lines = [
        ("项目", project_name or "-"),
        ("产物版本", " · ".join(f"{STAGE_LABELS.get(k, k)} v{v}"
                              for k, v in versions.items() if v) or "-"),
        ("需求总数", summary.get("fr_total") or 0),
        ("其中 P0", p0),
    ]
    for key in trace.DOWNSTREAM:
        miss = gaps.get(key) or []
        lines.append((f"P0 覆盖 · {STAGE_LABELS[key]}",
                      f"{covered.get(key, 0)}/{p0}"
                      + (f"（未覆盖：{'、'.join(miss)}）" if miss else "")))
    orphans = summary.get("orphan_testcases") or []
    lines.append(("未关联需求的用例", "、".join(orphans) if orphans else "无"))
    lines.append(("链路状态分布", " · ".join(
        f"{trace.STATUS_TEXT.get(k, k)} {status_count[k]}"
        for k in STATUS_ORDER if status_count.get(k)) or "-"))
    lines.append(("说明", "「未记录」表示对应产物生成于追溯能力上线之前，没有关联字段，"
                        "不计为断链；重跑该阶段即可补全链路。"))

    ws = wb.create_sheet("覆盖概览")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 78
    title = ws.cell(row=1, column=1,
                    value=f"{project_name} 需求追溯矩阵" if project_name else "需求追溯矩阵")
    title.font = Font(bold=True, size=13)
    for r, (k, v) in enumerate(lines, 3):
        kc = ws.cell(row=r, column=1, value=k)
        kc.font = Font(bold=True)
        kc.alignment = Alignment(vertical="top")
        vc = ws.cell(row=r, column=2, value=v)
        vc.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def export_trace_excel(data: dict, path: str, project_name: str = "") -> str:
    """把 trace_service.build_project_matrix 的结果写成 xlsx，返回文件路径。"""
    wb = Workbook()
    _matrix_sheet(wb, data.get("rows") or [])
    _overview_sheet(wb, data, project_name)
    wb.save(path)
    return path
