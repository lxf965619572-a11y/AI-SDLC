"""把阶段产物的 Markdown 文档渲染为 Word(.docx)。
支持：标题、表格、有序/无序列表、引用、代码块（含 ```mermaid / ```sql / ```json）、
行内加粗/斜体/行内代码。Mermaid 图以源码代码块保留（Word 中无法直接渲染矢量图）。"""
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*.+?\*)")
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}.*\|?\s*$")


def _set_mono(run, size=9):
    run.font.name = "Consolas"
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)


def _shade_paragraph(p, fill="F3F4F6"):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def _add_inline(paragraph, text: str):
    """解析行内 **bold** / `code` / *italic*，逐段写入 runs。"""
    for tok in _INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            paragraph.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            _set_mono(paragraph.add_run(tok[1:-1]))
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            paragraph.add_run(tok[1:-1]).italic = True
        else:
            paragraph.add_run(tok)


def _split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _flush_table(doc, rows: list[list[str]]):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    try:
        table.style = "Table Grid"
    except Exception:
        pass
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = table.cell(i, j)
            val = row[j] if j < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            _add_inline(p, val)
            if i == 0:
                for run in p.runs:
                    run.bold = True


def export_docx(markdown: str, path: str, title: str | None = None):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(10.5)

    lines = markdown.splitlines()
    i = 0
    table_buf: list[list[str]] = []
    in_code = False
    code_lang = ""
    code_lines: list[str] = []

    def flush_code():
        nonlocal code_lines, code_lang
        if code_lang == "mermaid":
            p = doc.add_paragraph()
            r = p.add_run("图（Mermaid 源码）")
            r.bold = True
            r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        for cl in code_lines:
            cp = doc.add_paragraph()
            _set_mono(cp.add_run(cl if cl else " "))
            _shade_paragraph(cp)
        code_lines = []
        code_lang = ""

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 代码块围栏
        if stripped.startswith("```"):
            if not in_code:
                if table_buf:
                    _flush_table(doc, table_buf)
                    table_buf = []
                in_code = True
                code_lang = stripped[3:].strip().lower()
            else:
                flush_code()
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # 表格行
        if stripped.startswith("|") and stripped.count("|") >= 2:
            if not _TABLE_SEP.match(stripped):
                table_buf.append(_split_row(stripped))
            i += 1
            continue
        elif table_buf:
            _flush_table(doc, table_buf)
            table_buf = []

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if m:
            level = min(len(m.group(1)), 4)
            h = doc.add_heading("", level=level)
            _add_inline(h, m.group(2))
            i += 1
            continue

        # 无序 / 有序列表
        lm = re.match(r"^[-*+]\s+(.*)", stripped)
        if lm:
            _add_inline(doc.add_paragraph(style="List Bullet"), lm.group(1))
            i += 1
            continue
        nm = re.match(r"^\d+[.)]\s+(.*)", stripped)
        if nm:
            _add_inline(doc.add_paragraph(style="List Number"), nm.group(1))
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            r = p.add_run(stripped.lstrip("> ").strip())
            r.italic = True
            r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            i += 1
            continue

        # 分隔线 / 空行
        if re.match(r"^(-{3,}|\*{3,})$", stripped) or not stripped:
            i += 1
            continue

        # 普通段落
        _add_inline(doc.add_paragraph(), stripped)
        i += 1

    if table_buf:
        _flush_table(doc, table_buf)
    if in_code and code_lines:
        flush_code()

    doc.save(path)
    return path
