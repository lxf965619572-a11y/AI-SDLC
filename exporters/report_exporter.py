"""软件测评报告装配（GJB 438B 风格，v1 取够用口径）。

第一性原理：测评报告的价值在于「每个结论都能回溯到一份产物或一条工具输出」。
所以本模块不调 LLM、不重新判断，只做两件事：
  1. assemble_report —— 把已落库的产物元数据折叠成一份结构化结论；
  2. report_markdown / export_report_* —— 把结论排成人能审、能归档的交付件。
凡是元数据里没有的，报告里就写「未产出」，绝不用推测补齐。
"""
from __future__ import annotations

from datetime import datetime

from core import trace

TITLE = "软件测评报告"
STANDARD = "GJB 438B《军用软件开发文档通用要求》"

# 报告里出现的阶段及其展示名（与 pipeline.nodes.STAGE_TITLES 口径一致，
# 这里只列要写进报告的那几个，避免报告被解析阶段的原始数据淹没）
REPORT_STAGES = ("requirement", "hld", "lld", "testcase",
                 "code", "static", "test_impl", "exec")
STAGE_LABELS = {"requirement": "软件需求规格说明", "hld": "概要设计说明",
                "lld": "详细设计说明", "testcase": "测试用例设计",
                "code": "代码实现", "static": "静态检查",
                "test_impl": "测试实现", "exec": "验证执行",
                "report": "软件测评报告", "parse": "结构化原始数据"}
STATUS_LABELS = {"approved": "已批准", "pending_review": "待人工裁决",
                 "rejected": "未通过"}
DECISION_LABELS = {"fix_code": "被测代码缺陷", "fix_test": "测试代码缺陷",
                   "blocked": "执行环境不可用", "next": "通过"}
VERDICT_LABELS = {"ok": "通过", "fail": "不通过", "skipped": "未执行"}


def _pct(value) -> str:
    return "-" if value is None else f"{float(value):.2f}%"


def _meta(arts: dict, stage: str) -> dict:
    return ((arts or {}).get(stage) or {}).get("meta") or {}


def _history(history: dict | None, arts: dict, stage: str) -> list[dict]:
    """某阶段的全部运行记录（升序）。没给 history 时退化为「只有最新一版」。"""
    runs = (history or {}).get(stage)
    if runs:
        return list(runs)
    art = (arts or {}).get(stage)
    return [art] if art else []


# ---------------- 各章节 ----------------
def _conclusion(arts: dict, static_m: dict, exec_m: dict) -> dict:
    """测评结论：静态与执行两个判据都通过才算通过，缺任一判据都不给通过结论。"""
    reasons = []
    if not exec_m:
        reasons.append("未产出验证执行结论（未配置验证机或流水线未跑到该阶段）")
    elif exec_m.get("skipped"):
        reasons.append(f"验证未执行：{exec_m.get('reason') or '执行环境不可用'}")
    elif not exec_m.get("ok"):
        reasons.append(f"验证执行未通过：{exec_m.get('reason') or '见测评结果'}")
    if not static_m:
        reasons.append("未产出静态检查结论")
    elif not static_m.get("ok"):
        reasons.append(f"静态检查存在必查项违规 {static_m.get('required', 0)} 条")
    tests = (exec_m.get("tests") or {})
    if tests and tests.get("consistent") is False:
        reasons.append("测试桩自报结果与实际输出不一致，全绿结论不可信")

    passed = not reasons
    if passed:
        cov = (exec_m.get("coverage") or {}).get("totals") or {}
        text = (f"被测软件通过全部测评项：{tests.get('passed')} 条用例全部通过，"
                f"分支覆盖 {_pct(cov.get('branch_pct'))}、"
                f"行覆盖 {_pct(cov.get('line_pct'))}，静态检查无必查项违规。")
    else:
        text = "被测软件未通过测评，理由见下；未闭环项须走问题报告单与偏差单处置。"
    return {"pass": passed, "text": text, "reasons": reasons}


def _environment(exec_m: dict) -> dict:
    env = exec_m.get("env") or {}
    return {"host": env.get("host") or env.get("uname") or "-",
            "uname": env.get("uname") or "-",
            "gcc": env.get("gcc") or "-",
            "gcov": env.get("gcov") or "-",
            "make": env.get("make") or "-",
            "cores": env.get("cores") or "-",
            "transport": exec_m.get("transport") or "-",
            "workdir": exec_m.get("workdir") or "-",
            "duration_s": exec_m.get("duration_s"),
            "commands": exec_m.get("commands") or [],
            "raw": env}


def _tests_section(exec_m: dict) -> dict:
    t = exec_m.get("tests") or {}
    return {"expected": t.get("expected") or [], "results": t.get("results") or [],
            "total": t.get("total") or 0, "passed": t.get("passed") or 0,
            "failed": t.get("failed") or 0, "missing": t.get("missing") or 0,
            "all_pass": bool(t.get("all_pass")),
            "consistent": t.get("consistent"),
            "crashed": bool(t.get("crashed")), "run_exit": t.get("run_exit"),
            "verdict": (exec_m.get("verdict") or {}).get("tests")}


def _coverage_section(exec_m: dict) -> dict:
    c = exec_m.get("coverage") or {}
    fns = [f for f in (c.get("functions") or []) if f.get("kind") != "file"]
    return {"totals": c.get("totals") or {}, "functions": fns,
            # 文件级结论单独存放在 coverage["files"]（summarize_coverage 已按 src/ 过滤）
            "files": c.get("files") or [],
            "below_branch": c.get("below_branch") or [],
            "below_line": c.get("below_line") or [],
            "untested_functions": c.get("untested_functions") or [],
            "thresholds": c.get("thresholds") or exec_m.get("thresholds") or {},
            "ok": bool(c.get("ok")),
            "verdict": (exec_m.get("verdict") or {}).get("coverage")}


def _static_section(static_m: dict) -> dict:
    return {"ok": bool(static_m.get("ok")),
            "total": static_m.get("total") or 0,
            "required": static_m.get("required") or 0,
            "advisory": static_m.get("advisory") or 0,
            "by_rule": static_m.get("by_rule") or {},
            "violations": static_m.get("violations") or [],
            "thresholds": static_m.get("thresholds") or {},
            "functions": static_m.get("functions") or [],
            "verdict": "ok" if static_m.get("ok") else ("fail" if static_m else "skipped")}


def _problems(history: dict | None, arts: dict) -> list[dict]:
    """问题报告单：把闭环过程中每一轮未通过的判据都留痕，包括后来修好的。

    「修好了」不等于「没发生过」。航天软件的问题归零要求能拿出问题清单与
    处置过程，所以历史轮次一律登记，并标注是否已闭环。"""
    out: list[dict] = []
    seq = 0
    for stage in ("static", "exec"):
        runs = _history(history, arts, stage)
        for i, run in enumerate(runs):
            meta = run.get("meta") or {}
            if not meta or meta.get("ok") or meta.get("skipped"):
                continue
            seq += 1
            later_ok = any((r.get("meta") or {}).get("ok")
                           for r in runs[i + 1:])
            att = meta.get("attribution") or {}
            if stage == "static":
                detail = (f"必查项违规 {meta.get('required', 0)} 条："
                          + "、".join(sorted({v.get("rule_id", "-")
                                              for v in meta.get("violations") or []
                                              if v.get("severity") == "required"})))
            else:
                detail = meta.get("reason") or "验证执行未通过"
            out.append({
                "id": f"PR-{seq:03d}", "stage": stage,
                "stage_label": STAGE_LABELS.get(stage, stage),
                "version": run.get("version"), "status": run.get("status"),
                "title": ("静态检查必查项违规" if stage == "static"
                          else "验证执行未通过"),
                "detail": detail,
                "decision": att.get("decision") or meta.get("decision") or "",
                "decision_label": DECISION_LABELS.get(
                    att.get("decision") or meta.get("decision") or "", "-"),
                "attribution": att.get("reason") or "",
                "attribution_source": att.get("source") or "",
                "resolved": bool(later_ok),
                "disposition": ("已闭环（后续轮次判据通过）" if later_ok
                                else "未闭环，转人工裁决"),
            })
    return out


def _deviations(arts: dict, static_m: dict) -> list[dict]:
    """偏差单：最终仍存在的必查项违规。可偏差项须人工批准，不可偏差项无权放行。"""
    if not static_m or static_m.get("ok"):
        return []
    approved = ((arts.get("static") or {}).get("status") == "approved")
    out = []
    for i, v in enumerate([x for x in static_m.get("violations") or []
                           if x.get("severity") == "required"], 1):
        out.append({"id": f"DEV-{i:03d}", "rule_id": v.get("rule_id", "-"),
                    "title": v.get("title", ""),
                    "standard_ref": v.get("standard_ref", ""),
                    "location": f"{v.get('file')}:{v.get('line') or '-'}",
                    "func": v.get("func") or "-",
                    "message": v.get("message", ""),
                    "waivable": bool(v.get("waivable")),
                    "disposition": ("人工批准放行" if (v.get("waivable") and approved)
                                    else "不可偏差，须整改" if not v.get("waivable")
                                    else "待人工裁决")})
    return out


def _evidence(exec_m: dict, static_m: dict) -> list[dict]:
    """证据清单：路径 + sha256。指纹在同步前于本机算出，证明跑的就是这份代码。"""
    out = []
    for path, info in sorted((exec_m.get("manifest") or {}).items()):
        # manifest 的值是 {"sha256", "bytes"}；兼容早期只存哈希字符串的产物
        sha = info.get("sha256") if isinstance(info, dict) else info
        nbytes = info.get("bytes") if isinstance(info, dict) else None
        out.append({"kind": "同步输入", "path": path, "sha256": sha or "-",
                    "bytes": nbytes})
    if exec_m.get("evidence"):
        out.append({"kind": "原始输出归档", "path": exec_m["evidence"],
                    "sha256": "见同目录 exec.json"})
    return out


def _trace_summary(arts: dict, coverage_min=None) -> dict:
    """需求追溯摘要：直接复用 core.trace，报告与前端矩阵不会各说一套。"""
    m = trace.build_matrix(
        _meta(arts, "requirement"), _meta(arts, "hld"), _meta(arts, "lld"),
        _meta(arts, "testcase"), code_meta=_meta(arts, "code") or None,
        static_meta=_meta(arts, "static") or None,
        test_impl_meta=_meta(arts, "test_impl") or None,
        exec_meta=_meta(arts, "exec") or None, coverage_min=coverage_min)
    s = m.get("summary") or {}
    return {"fr_total": s.get("fr_total") or 0, "p0_total": s.get("p0_total") or 0,
            "covered": s.get("covered") or {}, "uncovered_p0": s.get("uncovered_p0") or {},
            "status_counts": s.get("status_counts") or {},
            "orphan_testcases": s.get("orphan_testcases") or [],
            "rows": m.get("rows") or []}


def assemble_report(arts: dict, project_name: str = "", versions: dict | None = None,
                    history: dict | None = None, coverage_min=None) -> dict:
    """把各阶段产物折叠成一份结构化测评结论。

    arts: {stage: {"markdown", "meta", "version", "status"}}
    history: {stage: [同结构，按版本升序]}，用于问题报告单留痕；缺省只看最新版。"""
    arts = arts or {}
    static_m = _meta(arts, "static")
    exec_m = _meta(arts, "exec")
    versions = versions or {k: (v or {}).get("version") for k, v in arts.items()}
    return {
        "title": TITLE, "standard": STANDARD, "project_name": project_name or "-",
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "versions": versions,
        "artifacts": {st: {"version": (arts.get(st) or {}).get("version"),
                           "status": (arts.get(st) or {}).get("status"),
                           "status_label": STATUS_LABELS.get(
                               (arts.get(st) or {}).get("status"), "-"),
                           "chars": len((arts.get(st) or {}).get("markdown") or "")}
                      for st in REPORT_STAGES if st in arts},
        "conclusion": _conclusion(arts, static_m, exec_m),
        "environment": _environment(exec_m),
        "build": {"ok": (exec_m.get("build") or {}).get("ok"),
                  "verdict": (exec_m.get("verdict") or {}).get("build"),
                  "warning_count": (exec_m.get("build") or {}).get("warning_count") or 0,
                  "errors": (exec_m.get("build") or {}).get("errors") or []},
        "tests": _tests_section(exec_m),
        "coverage": _coverage_section(exec_m),
        "static": _static_section(static_m),
        "trace": _trace_summary(arts, coverage_min),
        "problems": _problems(history, arts),
        "deviations": _deviations(arts, static_m),
        "evidence": _evidence(exec_m, static_m),
        "thresholds": {"coverage_branch_min": coverage_min,
                       "complexity_max": (static_m.get("thresholds") or {}).get("complexity_max"),
                       "max_fix_rounds": None},
    }


# ---------------- Markdown 渲染 ----------------
def _table(headers: list[str], rows: list[list]) -> list[str]:
    out = ["| " + " | ".join(str(h) for h in headers) + " |",
           "|" + "|".join("---" for _ in headers) + "|"]
    for r in rows:
        cells = [str(c if c is not None else "-").replace("|", "/").replace("\n", " ")
                 for c in r]
        cells += [""] * (len(headers) - len(cells))
        out.append("| " + " | ".join(cells[:len(headers)]) + " |")
    return out


def report_markdown(data: dict) -> str:
    """渲染成 Markdown：既是前端展示的 report 阶段产物，也是 docx 导出的输入。"""
    c = data.get("conclusion") or {}
    env = data.get("environment") or {}
    tests = data.get("tests") or {}
    cov = data.get("coverage") or {}
    st = data.get("static") or {}
    tr = data.get("trace") or {}
    th = cov.get("thresholds") or {}
    totals = cov.get("totals") or {}
    L: list[str] = []

    L += [f"# {data.get('title') or TITLE}", "",
          f"> 编制依据：{data.get('standard') or STANDARD}　"
          f"测评对象：{data.get('project_name') or '-'}　"
          f"生成时间：{data.get('generated_at') or '-'}", "",
          "## 1 范围", "", "### 1.1 标识",
          "",
          ]
    L += _table(["文档", "版本"], [[STAGE_LABELS.get(k, k), f"v{v}"]
                                  for k, v in (data.get("versions") or {}).items() if v])
    L += ["", "### 1.2 系统概述", "",
          "本报告由代码验证闭环自动装配，覆盖静态检查、可执行测试与覆盖率采集三类判据；"
          "全部结论均来自确定性工具输出，不含模型推测。", "",
          "### 1.3 文档概述", "",
          "第 3 章给出测评环境与方法，第 4 章逐项给出判据结论，"
          "第 5 章为需求追溯摘要，第 6、7 章为问题报告单与偏差单，"
          "第 8 章为测评结论，第 9 章为证据清单。", ""]

    L += ["## 2 引用文档", "",
          "- " + (data.get("standard") or STANDARD),
          "- GJB 8114《C/C++语言编程安全子集》（条款号待标准化部门核定）",
          "- QJ 20084《航天嵌入式软件编码规范》（条款号待标准化部门核定）", ""]

    L += ["## 3 测评综述", "", "### 3.1 测评环境", ""]
    L += _table(["项", "值"], [
        ["执行环境", env.get("uname")], ["编译器", env.get("gcc")],
        ["覆盖率工具", env.get("gcov")], ["构建工具", env.get("make")],
        ["CPU 核数", env.get("cores")], ["传输方式", env.get("transport")],
        ["远端工作区", env.get("workdir")],
        ["本轮耗时(s)", env.get("duration_s")]])
    L += ["", "### 3.2 测评方法", "",
          "1. 静态检查：按受限子集规则表（与代码生成提示词同源）逐条判定，"
          "并核对详细设计函数是否全部落地；",
          "2. 可执行测试：在验证机上以 `gcc -std=c99 -Wall -Wextra` 编译并链接测试程序，"
          "运行后按 `TC-xxx PASS/FAIL` 逐条判定；",
          "3. 覆盖率：以 `-fprofile-arcs -ftest-coverage` 插桩，"
          "解析 `gcov -b -c` 输出得到函数级行/分支覆盖；",
          f"4. 判据门限：分支覆盖 ≥ {_pct(th.get('branch_min'))}，"
          f"行覆盖 ≥ {_pct(th.get('line_min'))}，圈复杂度 ≤ "
          f"{(st.get('thresholds') or {}).get('complexity_max', '-')}。", ""]

    L += ["## 4 测评结果", "", "### 4.1 构建", "",
          f"结论：**{VERDICT_LABELS.get((data.get('build') or {}).get('verdict'), '-')}**"
          f"　警告 {(data.get('build') or {}).get('warning_count', 0)} 条", ""]
    errs = (data.get("build") or {}).get("errors") or []
    if errs:
        L += _table(["文件", "行", "说明"],
                    [[e.get("file"), e.get("line"), e.get("msg")] for e in errs[:20]])
        L.append("")

    L += ["### 4.2 用例执行", "",
          f"结论：**{VERDICT_LABELS.get(tests.get('verdict'), '-')}**　"
          f"通过 {tests.get('passed', 0)} / 失败 {tests.get('failed', 0)} / "
          f"未执行 {tests.get('missing', 0)}（设计用例 {len(tests.get('expected') or [])} 条）",
          ""]
    if tests.get("results"):
        L += _table(["用例编号", "结果", "说明"],
                    [[r.get("id"), r.get("status_text") or r.get("status"),
                      r.get("detail")] for r in tests["results"]])
        L.append("")
    if tests.get("consistent") is False:
        L += ["> **警告**：测试桩自报的用例数与实际输出不一致，上表结论不可直接采信。", ""]
    if tests.get("crashed"):
        L += [f"> **警告**：测试进程异常退出（exit={tests.get('run_exit')}）。", ""]

    L += ["### 4.3 覆盖率", "",
          f"结论：**{VERDICT_LABELS.get(cov.get('verdict'), '-')}**　"
          f"行覆盖 {_pct(totals.get('line_pct'))}"
          f"（{totals.get('lines_hit', '-')}/{totals.get('lines_total', '-')}）　"
          f"分支覆盖 {_pct(totals.get('branch_pct'))}"
          f"（{totals.get('branch_taken', '-')}/{totals.get('branch_total', '-')}）", ""]
    fns = cov.get("functions") or []
    if fns:
        L += _table(["函数", "文件", "行覆盖", "分支覆盖", "分支数", "是否执行"],
                    [[f.get("name"), f.get("file"), _pct(f.get("lines_pct")),
                      _pct(f.get("branch_effective")), f.get("branch_total"),
                      "否" if f.get("never_executed") else "是"] for f in fns])
        L.append("")
    for label, key in (("未被执行的函数", "untested_functions"),
                       ("分支覆盖不足的函数", "below_branch"),
                       ("行覆盖不足的函数", "below_line")):
        if cov.get(key):
            L.append(f"- {label}：" + "、".join(cov[key]))
    if any(cov.get(k) for k in ("untested_functions", "below_branch", "below_line")):
        L.append("")

    L += ["### 4.4 静态检查", "",
          f"结论：**{'通过' if st.get('ok') else ('不通过' if st else '未产出')}**　"
          f"必查项 {st.get('required', 0)} 条 · 建议项 {st.get('advisory', 0)} 条", ""]
    if st.get("by_rule"):
        L += _table(["规则编号", "条数"],
                    [[k, v] for k, v in sorted(st["by_rule"].items())])
        L.append("")
    vs = st.get("violations") or []
    if vs:
        L += _table(["规则", "等级", "位置", "函数", "说明"],
                    [[v.get("rule_id"), "必查" if v.get("severity") == "required" else "建议",
                      f"{v.get('file')}:{v.get('line') or '-'}", v.get("func"),
                      v.get("message")] for v in vs[:60]])
        L.append("")

    L += ["## 5 需求追溯摘要", "",
          f"需求 {tr.get('fr_total', 0)} 条（其中 P0 {tr.get('p0_total', 0)} 条），"
          "行状态分布：" + ("、".join(
              f"{trace.STATUS_TEXT.get(k, k)} {v}"
              for k, v in (tr.get("status_counts") or {}).items()) or "-"), ""]
    gaps = tr.get("uncovered_p0") or {}
    gl = [f"{trace.LINK_LABELS.get(k, k)}：{'、'.join(v)}"
          for k, v in gaps.items() if v]
    if gl:
        L += ["P0 需求未覆盖项："] + [f"- {x}" for x in gl] + [""]
    if tr.get("orphan_testcases"):
        L += ["- 未关联需求的用例：" + "、".join(tr["orphan_testcases"]), ""]

    L += ["## 6 问题报告单", ""]
    probs = data.get("problems") or []
    if probs:
        L += _table(["编号", "阶段", "产物版本", "问题", "定责", "处置"],
                    [[p.get("id"), p.get("stage_label"), f"v{p.get('version')}",
                      p.get("detail"), p.get("decision_label"), p.get("disposition")]
                     for p in probs])
        L.append("")
        for p in probs:
            if p.get("attribution"):
                L.append(f"- {p['id']} 归因（{p.get('attribution_source') or '-'}）："
                         f"{p['attribution']}")
        L.append("")
    else:
        L += ["本轮测评未产生问题报告单。", ""]

    L += ["## 7 偏差单", ""]
    devs = data.get("deviations") or []
    if devs:
        L += _table(["编号", "规则", "可否偏差", "位置", "说明", "处置"],
                    [[d.get("id"), f"{d.get('rule_id')} {d.get('title')}",
                      "可" if d.get("waivable") else "不可", d.get("location"),
                      d.get("message"), d.get("disposition")] for d in devs])
        L += ["", "> 不可偏差项属安全性底线，人工亦无权放行，只能整改代码。", ""]
    else:
        L += ["无待处置偏差项。", ""]

    L += ["## 8 测评结论", "",
          f"**{'通过' if c.get('pass') else '未通过'}**　{c.get('text') or ''}", ""]
    for r in c.get("reasons") or []:
        L.append(f"- {r}")
    if c.get("reasons"):
        L.append("")
    arts_tbl = data.get("artifacts") or {}
    if arts_tbl:
        L += _table(["产物", "版本", "状态"],
                    [[STAGE_LABELS.get(k, k), f"v{v.get('version')}",
                      v.get("status_label")] for k, v in arts_tbl.items()])
        L.append("")

    L += ["## 9 证据清单", ""]
    ev = data.get("evidence") or []
    if ev:
        L += _table(["类型", "路径", "字节", "sha256"],
                    [[e.get("kind"), e.get("path"), e.get("bytes"), e.get("sha256")]
                     for e in ev])
        L += ["", "> 输入指纹在同步至验证机之前于本机计算，用以证明执行的就是本报告"
              "所述的这份代码与测试。", ""]
    else:
        L += ["未采集到执行证据（验证未执行）。", ""]
    return "\n".join(L)


# ---------------- 文件导出 ----------------
def export_report_docx(data: dict, path: str) -> str:
    from exporters.docx_exporter import export_docx
    return export_docx(report_markdown(data), path,
                       title=data.get("title") or TITLE)


def export_report_excel(data: dict, path: str) -> str:
    """报告附表（xlsx）：结论、用例、覆盖率、静态违规、问题与偏差、证据清单。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="4472C4")
    head_font = Font(color="FFFFFF", bold=True)

    def sheet(name, headers, widths, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        for col, (h, w) in enumerate(zip(headers, widths), 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(col)].width = w
        for r, row in enumerate(rows, 2):
            for col, v in enumerate(row, 1):
                cell = ws.cell(row=r, column=col,
                               value=v if v is not None else "-")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        return ws

    c = data.get("conclusion") or {}
    env = data.get("environment") or {}
    tests = data.get("tests") or {}
    cov = data.get("coverage") or {}
    totals = cov.get("totals") or {}
    st = data.get("static") or {}
    sheet("测评结论", ["项", "值"], [24, 84], [
        ["项目", data.get("project_name")],
        ["编制依据", data.get("standard")],
        ["生成时间", data.get("generated_at")],
        ["测评结论", "通过" if c.get("pass") else "未通过"],
        ["结论说明", c.get("text")],
        ["未通过理由", "\n".join(c.get("reasons") or []) or "无"],
        ["产物版本", " · ".join(f"{STAGE_LABELS.get(k, k)} v{v}"
                              for k, v in (data.get("versions") or {}).items() if v)],
        ["执行环境", env.get("uname")],
        ["编译器 / 覆盖率工具", f"{env.get('gcc')} / {env.get('gcov')}"],
        ["用例", f"通过 {tests.get('passed', 0)} / 失败 {tests.get('failed', 0)}"
               f" / 未执行 {tests.get('missing', 0)}"],
        ["覆盖率", f"行 {_pct(totals.get('line_pct'))} · 分支 {_pct(totals.get('branch_pct'))}"],
        ["静态检查", f"必查项 {st.get('required', 0)} 条 · 建议项 {st.get('advisory', 0)} 条"],
    ], first=True)

    sheet("用例结果", ["用例编号", "结果", "说明"], [14, 10, 80],
          [[r.get("id"), r.get("status_text") or r.get("status"), r.get("detail")]
           for r in tests.get("results") or []])
    sheet("覆盖率", ["函数", "文件", "行覆盖%", "分支覆盖%", "分支数", "是否执行"],
          [28, 20, 12, 12, 10, 10],
          [[f.get("name"), f.get("file"), f.get("lines_pct"),
            f.get("branch_effective"), f.get("branch_total"),
            "否" if f.get("never_executed") else "是"]
           for f in cov.get("functions") or []])
    sheet("静态检查", ["规则", "等级", "文件", "行", "函数", "说明"],
          [12, 8, 20, 8, 22, 70],
          [[v.get("rule_id"), "必查" if v.get("severity") == "required" else "建议",
            v.get("file"), v.get("line"), v.get("func"), v.get("message")]
           for v in st.get("violations") or []])
    probs = data.get("problems") or []
    devs = data.get("deviations") or []
    sheet("问题与偏差", ["类型", "编号", "阶段/规则", "说明", "定责/可否偏差", "处置"],
          [10, 12, 18, 60, 18, 26],
          [[p.get("title"), p.get("id"), p.get("stage_label"), p.get("detail"),
            p.get("decision_label"), p.get("disposition")] for p in probs]
          + [["偏差单", d.get("id"), f"{d.get('rule_id')} {d.get('title')}",
              f"{d.get('location')} {d.get('message')}",
              "可偏差" if d.get("waivable") else "不可偏差", d.get("disposition")]
             for d in devs])
    sheet("证据清单", ["类型", "路径", "字节", "sha256"], [16, 46, 10, 68],
          [[e.get("kind"), e.get("path"), e.get("bytes"), e.get("sha256")]
           for e in data.get("evidence") or []])
    wb.save(path)
    return path
