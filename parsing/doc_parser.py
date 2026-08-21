"""文档解析：docx / pdf / md / txt → 带标题层级的段落列表。"""
from pathlib import Path


def parse_document(path: str | Path, file_type: str) -> list[dict]:
    """返回 [{"level": int(0=正文,1-3=标题), "text": str}] 段落序列。"""
    path = Path(path)
    ft = file_type.lower()
    if ft == "docx":
        return _parse_docx(path)
    if ft == "pdf":
        return _parse_pdf(path)
    if ft in ("md", "markdown"):
        return _parse_markdown(path.read_text(encoding="utf-8", errors="ignore"))
    if ft == "txt":
        return [{"level": 0, "text": path.read_text(encoding="utf-8", errors="ignore")}]
    raise ValueError(f"不支持的文件类型: {file_type}")


def _parse_docx(path: Path) -> list[dict]:
    from docx import Document
    doc = Document(str(path))
    paras = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        level = 0
        style = (p.style.name or "").lower()
        if style.startswith("heading"):
            try:
                level = int(style.replace("heading", "").strip() or "1")
                level = min(max(level, 1), 3)
            except ValueError:
                level = 1
        paras.append({"level": level, "text": text})
    # 表格内容按正文补充
    for tbl in doc.tables:
        rows = []
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            paras.append({"level": 0, "text": "\n".join(rows)})
    return paras


def _parse_pdf(path: Path) -> list[dict]:
    import pdfplumber
    paras = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            for line in text.splitlines():
                line = line.strip()
                if line:
                    paras.append({"level": 0, "text": line})
    return paras


def _parse_markdown(text: str) -> list[dict]:
    paras = []
    buf = []

    def flush():
        if buf:
            paras.append({"level": 0, "text": "\n".join(buf).strip()})
            buf.clear()

    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            flush()
            hashes = len(stripped) - len(stripped.lstrip("#"))
            paras.append({"level": min(hashes, 3), "text": stripped.lstrip("#").strip()})
        elif stripped:
            buf.append(stripped)
        else:
            flush()
    flush()
    return paras
